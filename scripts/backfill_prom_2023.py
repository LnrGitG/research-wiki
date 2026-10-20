#!/usr/bin/env python3
"""Бэкфилл 2023 года промноменклатуры: сводный файл Proizvodstvo_mes_2023 → канон.

Зачем: месячные бюллетени `Prom_MM_YYYY.xlsx` за 2023 год Росстат с медиабанка снял,
но остался сводный годовой файл «Производство продукции … по месяцам»
(`Proizvodstvo_mes_2023.xlsx`): строки — продукты с кодом ОКПД2, столбцы — 12 месяцев.
Он закрывает сентябрь–декабрь 2023, которых иначе не хватает в 36-месячном окне.

Особенность сопоставления: в сводном файле единица измерения вписана в название через
запятую («Паркет щитовой в сборе,тыс.кв.м»), поэтому имя канонизируется в два шага —
по коду ОКПД2 (карта из текущих бюллетеней) и по названию без хвоста единицы.

Запуск:
    python3 scripts/backfill_prom_2023.py --simulate
    python3 scripts/backfill_prom_2023.py
"""
import argparse
import datetime as dt
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from direct_core import Writer, norm                      # noqa: E402
from backfill_prom_history import PREFIX, code_map_from_current  # noqa: E402

SRC = '/home/lnr/yc-wiki/raw/rosstat/operational/Proizvodstvo_mes_2023.xlsx'


def parse_consolidated(path):
    """[(продукт, код, [значения по месяцам 1..12])] со всех листов ОКВЭД."""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out = []
    for sheet in wb.sheetnames:
        if sheet == 'Содержание':
            continue
        rows = list(wb[sheet].iter_rows(values_only=True))
        hdr = None
        for i, r in enumerate(rows[:14]):
            if any(isinstance(c, str) and 'Январ' in c for c in r):
                hdr = i
                break
        if hdr is None:
            continue
        for r in rows[hdr + 1:]:
            name = r[0] if r else None
            code = r[1] if len(r) > 1 else None
            if not isinstance(name, str) or len(name.strip()) <= 3:
                continue
            vals = [v if isinstance(v, (int, float)) else None for v in r[2:14]]
            if any(v is not None for v in vals):
                out.append((name.strip(), str(code).strip() if code else None, vals))
    wb.close()
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--src', default=SRC)
    ap.add_argument('--from', dest='dfrom', default='2023-09')
    ap.add_argument('--to', dest='dto', default='2023-12')
    ap.add_argument('--simulate', action='store_true')
    ap.add_argument('--repo', default=os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    args = ap.parse_args()
    lo = dt.date(*map(int, args.dfrom.split('-')), 1)
    hi = dt.date(*map(int, args.dto.split('-')), 1)

    w = Writer(tag='backfill_prom', simulate=args.simulate)
    q = w.cur
    q.execute("""SELECT name_ru FROM catalog.metric WHERE name_ru LIKE 'Промпродукция, месяц%'""")
    known = {norm(r[0]): r[0] for r in q.fetchall()}
    by_code = code_map_from_current(args.repo, w)
    print('канонических имён: %d, кодов в карте: %d' % (len(known), len(by_code)), flush=True)

    items = parse_consolidated(args.src)
    print('позиций в сводном файле 2023: %d' % len(items), flush=True)

    stat = {'by_code': 0, 'by_name': 0, 'skipped': 0}
    total = 0
    for month in range(1, 13):
        ps = dt.date(2023, month, 1)
        if not (lo <= ps <= hi):
            continue
        label = 'rosstat_Proizvodstvo_mes_2023_%02d' % month
        rows = []
        for product, code, vals in items:
            value = vals[month - 1] if month - 1 < len(vals) else None
            if value is None:
                continue
            name = by_code.get(code) if code else None
            if name and norm(name) in known:
                stat['by_code'] += 1
            else:
                stripped = re.sub(r',\s*[^,]{1,40}$', '', product)
                cand = (PREFIX + stripped)[:200]
                name = known.get(norm(cand))
                if name:
                    stat['by_name'] += 1
                else:
                    stat['skipped'] += 1
                    continue
            rows.append(dict(metric_name=name, value=value, ps=(2023, month), region=w.RU,
                             sub='product: %s' % product[:150], freq='M', source='rosstat',
                             release=label, unit='unknown'))
        prep, ins, conf, nm = w.ingest(rows, group=label, verbose=False)
        total += ins
        print('2023-%02d: подготовлено %5d | вставлено %5d | конфликтов %4d | новых метрик %d'
              % (month, prep, ins, conf, nm), flush=True)
    print('сопоставление: по коду %d, по имени %d, пропущено %d (на строку)'
          % (stat['by_code'], stat['by_name'], stat['skipped']), flush=True)
    print('ИТОГО вставлено: %d' % total, flush=True)
    w.close()


if __name__ == '__main__':
    main()

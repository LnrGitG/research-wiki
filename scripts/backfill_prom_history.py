#!/usr/bin/env python3
"""Бэкфилл истории промноменклатуры: месячные бюллетени Росстата → канон.

Зачем: прямой загрузчик `direct_rosstat_prom.py` работает с июля 2026 и накопил по
два наблюдения на продукт. Ряд из двух точек непригоден ни для сезонной
корректировки, ни для лаговых регрессий, поэтому история добирается архивными
бюллетенями «Произведено важнейших видов продукции за <месяц> <год>».

Три отличия от прямого загрузчика:
  1. период берётся из шапки бюллетеня (лист «Содержание», «…за январь 2024 года»),
     а не из имени файла — архивные файлы названы по-разному (`prom_01-2024.xlsx`);
  2. грузится только вариант «месяц»: «пред. месяц» депрекейчен 20.10.2026, а
     «накопленным итогом» восстановимо из месячных;
  3. имя продукта канонизируется по существующему каталогу (точное совпадение, затем
     по коду ОКПД2 из графы B) — иначе переименования Росстата плодили бы
     метрики-двойники. Продукты, которых в каталоге нет, пропускаются и считаются.

Запуск:
    python3 scripts/backfill_prom_history.py --dir /home/lnr/prom_archive --simulate
    python3 scripts/backfill_prom_history.py --dir /home/lnr/prom_archive \
        --from 2023-09 --to 2026-08
"""
import argparse
import datetime as dt
import glob
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from direct_core import Writer, norm                      # noqa: E402
from direct_rosstat_prom import CODE_RE, RF_PREFIX, SKIP_SHEETS, num  # noqa: E402

MONTHS = {'январь': 1, 'февраль': 2, 'март': 3, 'апрель': 4, 'май': 5, 'июнь': 6,
          'июль': 7, 'август': 8, 'сентябрь': 9, 'октябрь': 10, 'ноябрь': 11,
          'декабрь': 12}
PERIOD_RE = re.compile(r'за\s+([а-яё]+)\s+(\d{4})', re.I)
PREFIX = 'Промпродукция, месяц: '


def period_of(wb, path):
    """(год, месяц) из шапки бюллетеня; резерв — из имени файла."""
    for name in (['Содержание'] if 'Содержание' in wb.sheetnames else wb.sheetnames[:1]):
        for row in wb[name].iter_rows(min_row=1, max_row=6, values_only=True):
            for cell in row:
                if isinstance(cell, str):
                    m = PERIOD_RE.search(cell)
                    if m and m.group(1).lower() in MONTHS:
                        return int(m.group(2)), MONTHS[m.group(1).lower()]
    base = os.path.basename(path).lower()
    m = re.search(r'(\d{2})[-_](\d{4})', base)
    return (int(m.group(2)), int(m.group(1))) if m else None


def period_of_file(path):
    """Дешёвая проверка периода: читается только шапка файла.

    Полный разбор бюллетеня занимает ~35 с; повторные прогоны по тому же каталогу
    не должны тратить это время на месяцы, которые уже в базе."""
    import openpyxl
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        ps = period_of(wb, path)
        wb.close()
        return ps
    except Exception:                              # noqa: BLE001 — битый архив не должен ронять прогон
        return None


def parse_bulletin(path):
    """[(продукт, код ОКПД2, единица, значение месяца)] по строкам РФ."""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ps = period_of(wb, path)
    out = []
    if ps:
        for sheet in wb.sheetnames:
            if sheet in SKIP_SHEETS:
                continue
            ws = wb[sheet]
            product = code = unit = None
            for row in ws.iter_rows(values_only=True):
                a = row[0] if len(row) > 0 else None
                b = row[1] if len(row) > 1 else None
                if isinstance(a, str) and isinstance(b, str) and CODE_RE.match(b.strip()) and len(a.strip()) > 3:
                    product, code, unit = a.strip(), b.strip(), None
                    continue
                if product and unit is None and isinstance(a, str) and a.strip():
                    bs = str(b).strip() if b is not None else ''
                    if bs.isdigit() and len(bs) <= 6:
                        unit = a.strip()
                        continue
                if isinstance(a, str) and a.strip().startswith(RF_PREFIX) and product:
                    vals = [num(row[i]) if len(row) > i else None for i in (2, 3, 4)]
                    if any(v is not None for v in vals):
                        out.append((product, code, unit, vals[0]))
                    product = code = unit = None
    wb.close()
    return ps, out


def code_map_from_current(repo, writer):
    """код ОКПД2 → каноническое имя: по текущим бюллетеням raw/prom/Prom_MM_YYYY.xlsx."""
    mapping = {}
    for path in sorted(glob.glob(os.path.join(repo, 'raw', 'prom', 'Prom_*.xlsx'))):
        _ps, rows = parse_bulletin(path)
        for product, code, _u, _v in rows:
            if code:
                # имя в каталоге хранится обрезанным до 200 символов
                # (direct_core.ensure_metric), поэтому канонизируем так же
                mapping.setdefault(code, (PREFIX + product)[:200])
    return mapping


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--dir', default='/home/lnr/prom_archive')
    ap.add_argument('--from', dest='dfrom', default='2023-09')
    ap.add_argument('--to', dest='dto', default='2026-08')
    ap.add_argument('--simulate', action='store_true')
    ap.add_argument('--repo', default=os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    args = ap.parse_args()
    lo = dt.date(*map(int, args.dfrom.split('-')), 1)
    hi = dt.date(*map(int, args.dto.split('-')), 1)

    files = sorted(f for f in glob.glob(os.path.join(args.dir, '*.xlsx'))
                   if os.path.basename(f).lower().startswith('prom'))
    if not files:
        sys.exit('нет файлов в %s' % args.dir)

    w = Writer(tag='backfill_prom', simulate=args.simulate)
    q = w.cur
    q.execute("""SELECT metric_code, name_ru FROM catalog.metric
                 WHERE name_ru LIKE 'Промпродукция, месяц%'""")
    known = {norm(r[1]): r[1] for r in q.fetchall()}
    print('канонических имён в каталоге: %d' % len(known), flush=True)
    by_code = code_map_from_current(args.repo, w)

    done = set()
    if not args.simulate:
        q.execute("""SELECT DISTINCT o.period_start FROM data.observation o
                     JOIN catalog.metric m USING(metric_id)
                     WHERE m.name_ru LIKE 'Промпродукция, месяц%'""")
        done = {r[0] for r in q.fetchall()}
        print('месяцев уже в базе: %d' % len(done), flush=True)

    total_ins = 0
    for f in files:
        ps = period_of_file(f)                     # дешёвая проверка до полного разбора
        if ps:
            period = dt.date(ps[0], ps[1], 1)
            if not (lo <= period <= hi):
                print('%s → %04d-%02d: вне окна, пропуск' % (os.path.basename(f), ps[0], ps[1]), flush=True)
                continue
            if period in done:
                print('%s → %04d-%02d: уже в базе' % (os.path.basename(f), ps[0], ps[1]), flush=True)
                continue
        try:
            ps, parsed = parse_bulletin(f)
        except Exception as e:                     # noqa: BLE001 — архивный файл может быть битым
            print('ОШИБКА разбора %s: %s' % (os.path.basename(f), e), flush=True)
            continue
        if not ps:
            print('ПРОПУСК (нет периода): %s' % os.path.basename(f), flush=True)
            continue
        year, month = ps
        period = dt.date(year, month, 1)
        if not (lo <= period <= hi):
            print('%s → %04d-%02d: вне окна, пропуск' % (os.path.basename(f), year, month), flush=True)
            continue
        if period in done:
            print('%s → %04d-%02d: уже в базе' % (os.path.basename(f), year, month), flush=True)
            continue
        label = 'rosstat_Prom_%02d_%04d' % (month, year)
        rows, skipped = [], 0
        for product, code, unit, value in parsed:
            # сравнение по обрезанному имени: каталог хранит name_ru[:200]
            name = known.get(norm((PREFIX + product)[:200])) or by_code.get(code)
            if not name or norm(name) not in known:
                skipped += 1
                continue
            sub = ('product: %s' % product) + (' | ед.: %s' % str(unit)[:40]
                                               if unit and str(unit).strip() not in ('', 'None', 'ND') else '')
            rows.append(dict(metric_name=name, value=value, ps=(year, month), region=w.RU,
                             sub=sub, freq='M', source='rosstat', release=label, unit='unknown'))
        prep, ins, conf, nm = w.ingest(rows, group=label, verbose=False)
        total_ins += ins
        print('%-22s → %04d-%02d: разобрано %5d | в канон %5d | вне каталога %4d | вставлено %5d | новых метрик %d'
              % (os.path.basename(f), year, month, len(parsed), prep, skipped, ins, nm), flush=True)
    print('ИТОГО вставлено: %d' % total_ins, flush=True)
    w.close()


if __name__ == '__main__':
    main()

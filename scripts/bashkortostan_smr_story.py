#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Сюжет по объёму строительных работ в Республике Башкортостан (СМР).

Собирает две стороны:
  * длинный ряд — годовые объёмы работ по ВЭД «Строительство» по субъектам РФ из
    базы (метрика y477110251, ЕМИСС/Росстат, млн руб., 2001-2024), с долей
    региона в сумме по субъектам;
  * свежую сторону — месячные данные Росстата из регионального мониторинга
    (`raw/rosstat/socioeconomic_regions/info-stat-07-2026/…/03-01 объем работ
    выполненных по ВЭД Строительство.xlsx`): накопленные объёмы 2025 и 2026 годов
    и индексы физического объёма к соответствующему периоду.

Печатает числа для записки и рисует две панели в
`visualizations/bashkortostan-smr-2026m07.png`.

Запуск: python3 scripts/bashkortostan_smr_story.py
"""
import sys
from pathlib import Path

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import openpyxl

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / 'scripts'))
from db_tunnel import query  # noqa: E402

REGION_ID = 54                    # core.region: 'bashkortostan'
ANNUAL_METRIC = 'y477110251'      # объём работ по ВЭД «Строительство», млн руб., годовая
XLSX = (REPO / 'raw/rosstat/socioeconomic_regions/info-stat-07-2026/info-stat-07-2026'
        / '03 строительство/03-01 объем работ выполненных по ВЭД Строительство.xlsx')
OUT_PNG = REPO / 'visualizations/bashkortostan-smr-2026m07.png'

MONTHS = ['январь', 'январь-февраль', 'январь-март', 'январь-апрель', 'январь-май',
          'январь-июнь', 'январь-июль', 'январь-август', 'январь-сентябрь',
          'январь-октябрь', 'январь-ноябрь', 'январь-декабрь']
# В исходнике подписи периодов иногда переносятся по строке («январь-\nмай »),
# поэтому ключи приводим к одному виду — без пробелов внутри.
PLACEHOLDERS = (-99999999, -77777777)


def norm(text):
    return ''.join(str(text).split())


def rows_of(sheet_name):
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb[sheet_name]
    return list(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True))


def extract(sheet_name, name_part):
    """{(год, период): значение} по строке региона."""
    rows = rows_of(sheet_name)
    year_row, period_row = rows[2], rows[3]
    target = [r for r in rows if r[0] and name_part in str(r[0])]
    if not target:
        raise SystemExit(f'строка «{name_part}» не найдена на листе «{sheet_name}»')
    row = target[0]
    out, year = {}, None
    for col in range(1, len(year_row)):
        if year_row[col]:
            year = str(year_row[col]).replace(' год', '').strip()
        period = period_row[col]
        if not period or year is None:
            continue
        out[(year, norm(period))] = row[col]
    return out


def annual_from_db():
    """Годовые объёмы по региону и по всем субъектам: [(год, млн руб, доля %)]."""
    region = {r[0].year: float(r[1]) for r in query(
        "SELECT period_start, value FROM core.observation_v2 o JOIN core.metric m USING(metric_id) "
        "WHERE m.metric_code=%s AND o.region_id=%s AND o.value > 0 "
        "AND o.value NOT IN (-99999999, -77777777) ORDER BY period_start",
        (ANNUAL_METRIC, REGION_ID))}
    total = {r[0].year: float(r[1]) for r in query(
        "SELECT o.period_start, sum(o.value) FROM core.observation_v2 o JOIN core.metric m USING(metric_id) "
        "JOIN core.region r ON r.region_id=o.region_id "
        "WHERE m.metric_code=%s AND r.level='region' AND o.value > 0 "
        "AND o.value NOT IN (-99999999, -77777777) GROUP BY 1", (ANNUAL_METRIC,))}
    out = []
    for year in sorted(region):
        share = 100 * region[year] / total[year] if total.get(year) else None
        out.append((year, region[year], share))
    return out


def rank_in(year):
    """Место региона среди субъектов и число субъектов в выборке."""
    rows = query(
        "SELECT o.region_id, o.value FROM core.observation_v2 o JOIN core.metric m USING(metric_id) "
        "JOIN core.region r ON r.region_id=o.region_id "
        "WHERE m.metric_code=%s AND o.period_start=%s AND r.level='region' ORDER BY o.value DESC",
        (ANNUAL_METRIC, f'{year}-01-01'))
    rows = [r for r in rows if r[1] and r[1] > 0]
    for i, (rid, _v) in enumerate(rows, 1):
        if rid == REGION_ID:
            return i, len(rows)
    return None, len(rows)


def main():
    annual = annual_from_db()
    levels = {y: v for y, v, _s in annual}
    shares = {y: s for y, _v, s in annual}

    cum = extract('Млн.рублей', 'Башкортост')
    ifo_period = extract('к соотв. периоду', 'Башкортост')
    ifo_period_rf = extract('к соотв. периоду', 'Российская')
    ifo_month = extract(' к соотв. месяцу', 'Башкортост')

    print('=== годовой ряд (млн руб. и доля в сумме по субъектам):')
    for y, v, s in annual:
        if y in (2001, 2010, 2015, 2019, 2020, 2021, 2022, 2023, 2024):
            print(f'  {y}: {v:,.1f} млн руб. ({v/1000:,.1f} млрд), доля {s:.2f}%')

    y2024 = levels[2024]
    print(f'\n2024 год: {y2024/1000:,.1f} млрд руб., доля {shares[2024]:.2f}% от суммы по субъектам')
    place, n_reg = rank_in(2024)
    print(f'место среди субъектов в 2024: {place} из {n_reg}')

    print('\n=== накопленные объёмы 2025 и 2026 (млн руб.):')
    for m in MONTHS[:7]:
        v25 = cum.get(('2025', m))
        v26 = cum.get(('2026', m))
        if v25 and v26:
            print(f'  {m:16} 2025 {v25:>10,.1f} | 2026 {v26:>10,.1f} | номинал {100*v26/v25:5.1f}%')

    print('\n=== индекс физического объёма к соответствующему периоду (накопленно):')
    for m in MONTHS[:7]:
        b = ifo_period.get(('2026', m))
        r = ifo_period_rf.get(('2026', m))
        if b or r:
            print(f'  {m:16} Башкортостан {b} | РФ {r}')

    print('\n=== годовые ИФО (декабрь, к соответствующему периоду):')
    for y in ('2021', '2022', '2023', '2024', '2025'):
        print(f'  {y}: Башкортостан {ifo_period.get((y, "январь-декабрь"))} | '
              f'РФ {ifo_period_rf.get((y, "январь-декабрь"))}')

    print('\n=== месячные ИФО (к соответствующему месяцу 2025):')
    for m in MONTHS[:7]:
        v = ifo_month.get(('2026', m))
        if v:
            print(f'  {m:16} {v}')

    # ---- график ----
    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(15.5, 6.4))
    yrs = [y for y, _v, _s in annual]
    vals = [levels[y] / 1000 for y in yrs]
    ax1.bar(yrs, vals, color='#3b6ea5', width=0.72, label='объём работ, млрд руб.')
    ax1.set_title('Башкортостан: объём работ по ВЭД «Строительство»,\nгодовые значения', fontsize=12)
    ax1.set_ylabel('млрд руб. (текущие цены)')
    ax1.grid(axis='y', alpha=0.25)
    ax1b = ax1.twinx()
    ax1b.plot(yrs, [shares[y] for y in yrs], color='#c1440e', marker='o', ms=3.2, lw=1.6,
              label='доля в сумме по субъектам РФ, %')
    ax1b.set_ylabel('доля, %', color='#c1440e')
    ax1b.tick_params(axis='y', colors='#c1440e')
    h1, l1 = ax1.get_legend_handles_labels()
    h2, l2 = ax1b.get_legend_handles_labels()
    ax1.legend(h1 + h2, l1 + l2, loc='upper left', fontsize=9)

    labels = ['I', 'I-II', 'I-III', 'I-IV', 'I-V', 'I-VI', 'I-VII']
    months7 = MONTHS[:7]
    x = range(len(labels))
    v25 = [(cum.get(('2025', m)) or 0) / 1000 for m in months7]
    v26 = [(cum.get(('2026', m)) or 0) / 1000 for m in months7]
    ax2.bar([i - 0.2 for i in x], v25, width=0.38, color='#9db8d2', label='2025, млрд руб.')
    ax2.bar([i + 0.2 for i in x], v26, width=0.38, color='#3b6ea5', label='2026, млрд руб.')
    ax2.set_xticks(list(x))
    ax2.set_xticklabels(labels)
    ax2.set_ylabel('млрд руб. накопленно')
    ax2.set_title('Январь-июль: накопленные объёмы и отставание\nот прошлого года', fontsize=12)
    ax2.grid(axis='y', alpha=0.25)
    ax2b = ax2.twinx()
    b_path = [ifo_period.get(('2026', m)) for m in months7]
    r_path = [ifo_period_rf.get(('2026', m)) for m in months7]
    xs = [i for i, v in enumerate(b_path) if v]
    ax2b.plot([i for i, v in enumerate(b_path) if v], [b_path[i] for i in xs],
              color='#c1440e', marker='o', lw=2, label='ИФО Башкортостана, % г/г')
    ax2b.plot([i for i, v in enumerate(r_path) if v], [r_path[i] for i in xs],
              color='#5a8f29', marker='s', lw=1.6, ls='--', label='ИФО РФ, % г/г')
    ax2b.axhline(100, color='#666', lw=0.9, ls=':')
    ax2b.set_ylabel('ИФО к соответствующему периоду, %')
    h1, l1 = ax2.get_legend_handles_labels()
    h2, l2 = ax2b.get_legend_handles_labels()
    ax2.legend(h1 + h2, l1 + l2, loc='upper left', fontsize=9)

    fig.suptitle('Строительные работы в Республике Башкортостан: длинный ряд и 2026 год',
                 fontsize=14)
    fig.tight_layout(rect=(0, 0, 1, 0.95))
    OUT_PNG.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(OUT_PNG, dpi=150)
    print(f'\nграфик: {OUT_PNG}')


if __name__ == '__main__':
    main()
#!/usr/bin/env python3
"""Загрузчик макроэкономического опроса Банка России (профессиональные
прогнозисты, «макропрос») в канон research_wiki.

Источник: https://cbr.ru/statistics/ddkp/mo_br/ — на странице приложен один
файл `full.xlsx` со всей историей опросов (17 показателей, статистики по
респондентам, разрез прогнозных лет). Это единственный рациональный путь:
HTML-таблица на странице содержит только текущий выпуск, а история лежит
именно в приложении.

Ссылка на файл вида /Content/Document/File/<id>/full.xlsx отдаётся только с
Referer страницы — без него ЦБ возвращает 403.

Структура файла. Листы «1»…«17», по показателю на лист; лист «Навигация»
задаёт соответствие. В каждом листе шапка — строка с датами опросов, ниже
блоки статистик (медиана, среднее, макс, 90-й и 10-й процентили, квартили,
мин, число ответов). У показателей 1–15 в колонке «Прогнозный период» стоит
год прогноза, поэтому блок статистики занимает несколько строк — по строке
на прогнозный год. У 16–17 («нейтральная ставка», «долгосрочный рост ВВП»)
прогнозного года нет: одна строка на статистику.

В канон идут медиана, 10-й и 90-й процентили — медиана как центральная
ожидаемая траектория, перцентили как мера разброса прогнозов (прокси
неопределённости для риск-канала). Прогнозный год хранится в `sub_dimension`
(`target_2026`, `lt` для бессрочных показателей), дата опроса — в
`period_start`.

Полный разбор всех девяти статистик — в `data/cbr_mo_panel.csv`.

Запуск: --all (по умолчанию), --fetch, --parse, --load, --verify.
"""
import argparse
import csv
import datetime as dt
import hashlib
import json
import os
import re
import ssl
import subprocess
import sys
import unicodedata
import urllib.request

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PAGE = 'https://cbr.ru/statistics/ddkp/mo_br/'
RAW_DIR = os.path.join(REPO, 'raw/cbr/mo')
RAW_XLSX = os.path.join(RAW_DIR, 'full.xlsx')
META = os.path.join(RAW_DIR, 'meta.json')
SNAP = os.path.join(RAW_DIR, 'parse_snapshot.json')
PANEL = os.path.join(REPO, 'data/cbr_mo_panel.csv')
SOURCE_ID = 7
FREQ_M = 5
REGION_RU = 1

# slug, лист, название, единица (unit_code)
INDICATORS = [
    ('cpi_dec', '1', 'ИПЦ, % дек. к дек. предыдущего года', 'pct'),
    ('cpi_avg', '2', 'ИПЦ, % к предыдущему году в среднем за год', 'pct'),
    ('key_rate', '3', 'Ключевая ставка, % годовых в среднем за год', 'pct'),
    ('gdp', '4', 'ВВП, %, г/г', 'pct'),
    ('unemp_avg', '5', 'Уровень безработицы, % в среднем за год', 'pct'),
    ('unemp_dec', '6', 'Уровень безработицы, % дек. без исключения сезонности (до апреля 2025)', 'pct'),
    ('wage_nom', '7', 'Номинальная заработная плата, %, г/г в среднем за год', 'pct'),
    ('budget_bal', '8', 'Баланс консолидированного бюджета, % ВВП', 'pct'),
    ('exports', '9', 'Экспорт товаров и услуг, млрд долл. США в год', 'bln_usd'),
    ('imports', '10', 'Импорт товаров и услуг, млрд долл. США в год', 'bln_usd'),
    ('usd_rub', '11', 'Курс USD/RUB, руб. за долл. в среднем за год', 'rub_per_usd'),
    ('oil_tax', '12', 'Цена нефти для налогообложения, долл. США за баррель', 'usd_per_bbl'),
    ('oil_brent', '13', 'Цена нефти Brent, долл. США за баррель (до марта 2025)', 'usd_per_bbl'),
    ('oil_urals', '14', 'Цена нефти Urals, долл. США за баррель (до декабря 2023)', 'usd_per_bbl'),
    ('cds_5y', '15', 'CDS-спред по России 5Y, б.п. (до апреля 2022)', 'bp'),
    ('neutral_rate', '16', 'Нейтральная ключевая ставка, % годовых', 'pct'),
    ('gdp_lt', '17', 'Долгосрочный рост ВВП, %, г/г', 'pct'),
]

# подпись в файле -> код статистики
STATS = {
    'медиана': 'p50', 'среднее': 'mean', 'макс': 'max',
    '90-й процентиль': 'p90', '3-й квартиль': 'p75', '1-й квартиль': 'p25',
    '10-й процентиль': 'p10', 'мин': 'min', 'всего ответов': 'n',
}
IN_CANON = ('p50', 'p10', 'p90')
UNIT_NAMES = {
    'bln_usd': 'млрд долларов США',
    'rub_per_usd': 'руб. за доллар США',
    'usd_per_bbl': 'долл. США за баррель',
    'bp': 'базисный пункт',
}


def norm(s):
    s = unicodedata.normalize('NFKC', str(s)).lower().replace('ё', 'е')
    return re.sub(r'\s+', ' ', s).strip()


# -------------------------------------------------------------------- сырьё
def _curl(url, out_path, referer=None):
    """ЦБ отбивает urllib по TLS/заголовкам (403), curl проходит — берём его."""
    cmd = ['curl', '-sS', '-k', '-L', '--max-time', '120',
           '-A', ('Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 '
                  '(KHTML, like Gecko) Chrome/126.0 Safari/537.36'),
           '-H', 'Accept-Language: ru-RU,ru;q=0.9,en;q=0.8',
           '-o', out_path, '-w', '%{http_code}', url]
    if referer:
        cmd[1:1] = ['-H', 'Referer: ' + referer]
    code = subprocess.run(cmd, capture_output=True, text=True).stdout.strip()
    if code != '200':
        raise RuntimeError('curl вернул HTTP %s для %s' % (code, url))
    return code


def fetch():
    os.makedirs(RAW_DIR, exist_ok=True)
    page_html = os.path.join(RAW_DIR, '_page.html')
    _curl(PAGE, page_html)
    html = open(page_html, encoding='utf-8', errors='replace').read()
    os.remove(page_html)

    m = re.findall(r'/Content/Document/File/(\d+)/full\.xlsx', html)
    if not m:
        raise RuntimeError('на странице макропроса нет ссылки на full.xlsx')
    file_id = max(m, key=int)
    url = 'https://cbr.ru/Content/Document/File/%s/full.xlsx' % file_id

    tmp = RAW_XLSX + '.part'
    _curl(url, tmp, referer=PAGE)
    blob = open(tmp, 'rb').read()
    if blob[:2] != b'PK':
        os.remove(tmp)
        raise RuntimeError('ответ не xlsx (нет сигнатуры PK)')
    sha = hashlib.sha256(blob).hexdigest()
    old = json.load(open(META)) if os.path.exists(META) else None
    if old and old.get('sha256') == sha and os.path.exists(RAW_XLSX):
        os.remove(tmp)
        print('файл не изменился (sha256 %s…)' % sha[:12])
    else:
        os.replace(tmp, RAW_XLSX)
        print('скачано full.xlsx (%d байт, sha256 %s…)' % (len(blob), sha[:12]))
    json.dump({'url': url, 'page': PAGE, 'file_id': file_id, 'sha256': sha,
               'bytes': len(blob), 'fetched_at': dt.datetime.now().isoformat(timespec='seconds')},
              open(META, 'w'), ensure_ascii=False, indent=1)
    return RAW_XLSX


# ------------------------------------------------------------------ разбор
def parse(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    records, meta = [], {}
    for slug, sheet, name, unit in INDICATORS:
        if sheet not in wb.sheetnames:
            raise RuntimeError('нет листа %r (показатель %s)' % (sheet, slug))
        ws = wb[sheet]
        rows = list(ws.iter_rows(max_row=400, max_col=80, values_only=True))
        hdr_i, dates, single = None, {}, False
        for i, r in enumerate(rows[:12]):
            dts = [(j, c) for j, c in enumerate(r) if isinstance(c, dt.datetime)]
            if len(dts) >= 5:
                hdr_i = i
                if all(j >= 4 for j, _ in dts):
                    dates = {j: c.date() for j, c in dts}
                else:
                    dates = {j: c.date() for j, c in dts if j >= 3}
                    single = True
                break
        if hdr_i is None:
            raise RuntimeError('лист %s: не найдена шапка с датами опросов' % sheet)

        stat, seen_stats, n_rows = None, set(), 0
        for r in rows[hdr_i + 1:]:
            lab = norm(r[1]) if len(r) > 1 and r[1] else ''
            if lab in STATS:
                stat = STATS[lab]
                seen_stats.add(stat)
                target = 'lt' if single else None
            elif lab:
                continue
            elif stat is None:
                continue
            else:
                target = 'lt' if single else None
            if not single:
                t = r[3] if len(r) > 3 else None
                if not isinstance(t, dt.datetime):
                    continue
                target = 'target_%d' % t.year
            elif not single and target is None:
                continue
            for j, d in dates.items():
                v = r[j] if j < len(r) else None
                if isinstance(v, (int, float)):
                    records.append({'indicator': slug, 'stat': stat,
                                    'survey': d.isoformat(), 'target': target,
                                    'value': float(v)})
                    n_rows += 1
        missing = [k for k in ('p50', 'p10', 'p90') if k not in seen_stats]
        if missing:
            raise RuntimeError('лист %s (%s): нет статистик %s — формат изменился'
                               % (sheet, slug, ', '.join(missing)))
        meta[slug] = n_rows

    if not records:
        raise RuntimeError('разбор дал ноль записей')
    surveys = sorted({r['survey'] for r in records})
    return {'source_file': os.path.basename(path), 'records': len(records),
            'surveys': len(surveys), 'first_survey': surveys[0],
            'last_survey': surveys[-1], 'by_indicator': meta, 'rows': records}


# ------------------------------------------------------------------- запись
def ensure_unit(cur, code):
    cur.execute("select unit_id from core.unit where unit_code=%s", (code,))
    r = cur.fetchone()
    if r:
        return r[0]
    cur.execute("""insert into core.unit (unit_code, name_ru) values (%s, %s)
                   returning unit_id""", (code, UNIT_NAMES[code]))
    print('  заведена единица измерения: %s' % code)
    return cur.fetchone()[0]


def ensure_metric(cur, code, name, desc, unit_id):
    cur.execute("select metric_id from core.metric where metric_code=%s", (code,))
    r = cur.fetchone()
    if r:
        return r[0]
    cur.execute("""insert into core.metric (metric_code, name_ru, description, frequency_id,
                     unit_id, metric_type, is_derived, status)
                   values (%s, %s, %s, %s, %s, 'primary', false, 'active')
                   returning metric_id""", (code, name[:200], desc, FREQ_M, unit_id))
    return cur.fetchone()[0]


def ensure_release(cur, label, published):
    cur.execute("select release_id from core.release where release_label=%s", (label,))
    r = cur.fetchone()
    if r:
        cur.execute("update core.release set published_at=%s where release_id=%s",
                    (published, r[0]))
        return r[0]
    cur.execute("""insert into core.release (source_id, release_label, published_at, status)
                   values (%s, %s, %s, 'registered') returning release_id""",
                (SOURCE_ID, label, published))
    return cur.fetchone()[0]


def load(conn, snap):
    cur = conn.cursor()
    last = dt.date.fromisoformat(snap['last_survey'])
    label = 'cbr_mo_survey_%04d-%02d' % (last.year, last.month)
    nxt = dt.date(last.year + (last.month == 12), 1 if last.month == 12 else last.month + 1, 1)
    rel = ensure_release(cur, label, nxt - dt.timedelta(days=1))

    ids, units = {}, {}
    for slug, _s, name, unit in INDICATORS:
        units[unit] = ensure_unit(cur, unit)
        for st, human in (('p50', 'медиана'), ('p10', '10-й процентиль'),
                          ('p90', '90-й процентиль')):
            code = 'cbr_mo_%s_%s_m' % (slug, st)
            ids[(slug, st)] = ensure_metric(
                cur, code, '%s — %s прогнозов' % (name, human),
                'Макроэкономический опрос Банка России (профессиональные '
                'прогнозисты): %s, %s. Дата наблюдения — месяц опроса, '
                'прогнозный год — в sub_dimension (lt = долгосрочный ориентир).'
                % (name, human), units[unit])

    cur.execute("delete from core.observation_v2 where release_id=%s and metric_id = any(%s)",
                (rel, list(ids.values())))
    cleaned = cur.rowcount
    n = 0
    for r in snap['rows']:
        if r['stat'] not in IN_CANON:
            continue
        cur.execute("""insert into core.observation_v2
                         (metric_id, region_id, frequency_id, period_start, period_end, value,
                          assessment_type, observation_status, source_id, release_id, sub_dimension)
                       values (%s, %s, %s, %s, %s, %s, 'final', 'raw', %s, %s, %s)""",
                    (ids[(r['indicator'], r['stat'])], REGION_RU, FREQ_M,
                     r['survey'], r['survey'], r['value'], SOURCE_ID, rel, r['target']))
        n += 1
    cur.execute("update core.release set status='loaded' where release_id=%s", (rel,))
    conn.commit()
    print('релиз %s: снято %d, вставлено %d (медиана, p10, p90)' % (label, cleaned, n))
    return label


def write_panel(snap):
    os.makedirs(os.path.dirname(PANEL), exist_ok=True)
    with open(PANEL, 'w', newline='') as f:
        w = csv.writer(f)
        w.writerow(['indicator', 'statistic', 'survey_date', 'target', 'value'])
        for r in sorted(snap['rows'], key=lambda x: (x['indicator'], x['stat'], x['survey'], x['target'])):
            w.writerow([r['indicator'], r['stat'], r['survey'], r['target'], '%.6g' % r['value']])
    print('панель всех девяти статистик: %s (%d строк)' % (PANEL, len(snap['rows'])))


def verify(conn, snap):
    cur = conn.cursor()
    cur.execute("""select m.metric_code, o.period_start::text, o.sub_dimension, o.value
                   from data.observation o join catalog.metric m on m.metric_id=o.metric_id
                   where m.metric_code like 'cbr_mo\\_%%'""")
    canon = {}
    for c, p, sub, v in cur.fetchall():
        st = c.rsplit('_', 2)[-2]
        slug = c[len('cbr_mo_'):-len('_%s_m' % st)]
        canon[(slug, st, p, sub)] = float(v)
    src = {(r['indicator'], r['stat'], r['survey'], r['target']): r['value']
           for r in snap['rows'] if r['stat'] in IN_CANON}
    miss = [k for k in src if k not in canon]
    bad = [(k, src[k], canon[k]) for k in src if k in canon and abs(src[k] - canon[k]) > 1e-9]
    extra = [k for k in canon if k not in src]
    print('сверка: источник %d, канон %d, расхождений %d, пропущено %d, лишних %d'
          % (len(src), len(canon), len(bad), len(miss), len(extra)))
    for k, a, b in bad[:5]:
        print('  расхождение %s: %s против %s' % (k, a, b))
    return not miss and not bad and not extra


def main():
    ap = argparse.ArgumentParser()
    for f in ('fetch', 'parse', 'load', 'verify', 'all'):
        ap.add_argument('--' + f, action='store_true')
    a = ap.parse_args()
    if not any((a.fetch, a.parse, a.load, a.verify, a.all)):
        a.all = True

    path = RAW_XLSX
    if a.fetch or a.all:
        path = fetch()
    if a.parse or a.all:
        if not os.path.exists(path):
            raise SystemExit('нет сырья: сначала --fetch')
        snap = parse(path)
        json.dump(snap, open(SNAP, 'w'), ensure_ascii=False, indent=1)
        print('разобрано: %d записей, %d опросов, %s .. %s'
              % (snap['records'], snap['surveys'], snap['first_survey'], snap['last_survey']))
        write_panel(snap)
    else:
        snap = json.load(open(SNAP))

    import db_tunnel
    conn = db_tunnel.connect()
    if a.load or a.all:
        load(conn, snap)
    if a.verify or a.all:
        ok = verify(conn, snap)
        conn.close()
        sys.exit(0 if ok else 1)
    conn.close()


if __name__ == '__main__':
    main()

#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Загрузка факторов-ограничителей инвестиционной активности (ИИБ) из XLSX ЦБ в SQLite.

Файл: https://www.cbr.ru/Content/Document/File/194116/mp_survey_data_facts.xlsx
Лист 'Строительство': 7 факторов x 4 размерные группы, кварталы 1к20-2к26.
Идемпотентен.
"""
import urllib.request, ssl, io, re, sqlite3, os
from openpyxl import load_workbook

URL = 'https://www.cbr.ru/Content/Document/File/194116/mp_survey_data_facts.xlsx'
DB = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                  'data', 'rosstat_construction.db')
GROUPS = [(1, 'Экономика всего'), (9, 'Крупные'), (17, 'Средние'), (25, 'Малые и микро')]


def norm_period(p):
    m = re.match(r'(\d)к(\d\d)', str(p))
    if not m:
        return None
    return '20%sQ%s' % (m.group(2), m.group(1))


def main():
    ctx = ssl.create_default_context()
    ctx.check_hostname = False
    ctx.verify_mode = ssl.CERT_NONE
    req = urllib.request.Request(URL, headers={'User-Agent': 'Mozilla/5.0'})
    data = urllib.request.urlopen(req, timeout=90, context=ctx).read()
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    allf = list(wb['Строительство'].iter_rows(values_only=True))
    factors = [str(x) for x in allf[3][1:8]]
    con = sqlite3.connect(DB)
    cur = con.cursor()
    cur.execute('DROP TABLE IF EXISTS cbr_monitoring_invest_factors')
    cur.execute('CREATE TABLE cbr_monitoring_invest_factors ('
                'size_group TEXT, factor TEXT, period TEXT, value REAL, source TEXT)')
    rows = []
    for r in allf[4:]:
        per = norm_period(r[0]) if r[0] else None
        if not per:
            continue
        for gcol, gname in GROUPS:
            for fi, f in enumerate(factors):
                idx = gcol + fi
                v = r[idx] if idx < len(r) else None
                if isinstance(v, (int, float)):
                    rows.append((gname, f, per, float(v), 'cbr_mp_survey_facts'))
    cur.executemany('INSERT INTO cbr_monitoring_invest_factors VALUES (?,?,?,?,?)', rows)
    con.commit()
    print('строк: %d' % len(rows))
    for r in cur.execute('SELECT size_group, COUNT(*) FROM cbr_monitoring_invest_factors GROUP BY 1'):
        print(r)
    print('диапазон периодов:', cur.execute(
        'SELECT MIN(period), MAX(period) FROM cbr_monitoring_invest_factors').fetchone())
    con.close()


if __name__ == '__main__':
    main()
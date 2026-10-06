#!/usr/bin/env python3
"""Синхронизация региональных рядов ЦБ (формы 0420218) прямой закачкой.

Запускается ingest-watcher'ом на ВМ по задаче из ops.ingest_queue (ставится
pg_cron-джобом раз в месяц). Работает самодостаточно:

  1. Скачивает страницу «Статистика → Банковский сектор → Ипотека» и находит
     ссылки на XLSX-формы 02_04…02_15 (раздел vfs/statistics).
  2. Парсит листы 02_06/02_07/02_08/02_09/02_15 (региональный разрез:
     число ИЖК, объём, ставка, срок, задолженность, ДДУ).
  3. Пишет в data.observation с новым release-ом cbr-vfs-sync <дата>.

Идемпотентен: ON CONFLICT по (metric_id, region_id, period_start,
frequency_id, assessment_type, sub_dimension, release_id) — повторный
прогон того же релиза ничего не дублирует. Питфоллы (из transfer_cbr_vfs):
строчная нормализация имён регионов, алиасы без НАО/ХМАО-ЯНАО, агрегаты
РФ/ФО уже заведены в catalog.region со status source_aggregate.
"""
import io
import os
import re
import subprocess
import sys
from datetime import datetime

import openpyxl
import psycopg

PAGE = "https://www.cbr.ru/statistics/bank_sector/mortgage/"
SHEET_MAP = {
    ("02_06", "в рублях"): "cbr_izhk_count_m",
    ("02_06", "итого"): "cbr_izhk_count_m",
    ("02_07", "в рублях"): "cbr_izhk_volume_m",
    ("02_07", "итого"): "cbr_izhk_volume_m",
    ("02_08", "ставка в рублях"): "cbr_izhk_rate_m",
    ("02_08", "срок в руб"): "cbr_izhk_term_m",
    ("02_09", "в рублях"): "cbr_izhk_debt_m",
    ("02_09", "итого"): "cbr_izhk_debt_m",
    ("02_15", "в рублях"): "cbr_izhk_ddu_volume_m",
}
FORMS = {"02_04", "02_05", "02_06", "02_07", "02_08", "02_09", "02_15"}
MON = {"Январь": 1, "Февраль": 2, "Март": 3, "Апрель": 4, "Май": 5, "Июнь": 6,
       "Июль": 7, "Август": 8, "Сентябрь": 9, "Октябрь": 10, "Ноябрь": 11, "Декабрь": 12}
ALIAS = {
    "архангельская область": "архангельская область (без нао)",
    "кемеровская область - кузбасс": "кемеровская область — кузбасс",
    "тюменская область": "тюменская область (без хмао и янао)",
}
NAMES = {
    "cbr_izhk_count_m": ("ЦБ: число выданных ипотечных кредитов (ИЖК), шт, за месяц", 4),
    "cbr_izhk_volume_m": ("ЦБ: объём выдач ипотечных кредитов (ИЖК), млн руб., за месяц", 1),
    "cbr_izhk_rate_m": ("ЦБ: средневзвешенная ставка по ипотечным кредитам (руб.), %", 5),
    "cbr_izhk_term_m": ("ЦБ: средний срок кредитования по ИЖК (руб.), месяцев", 8),
    "cbr_izhk_debt_m": ("ЦБ: задолженность по ипотечным кредитам (руб.), млн руб., на начало месяца", 1),
    "cbr_izhk_ddu_volume_m": ("ЦБ: объём выдач ипотечных кредитов на приобретение жилья с ДДУ, млн руб., за месяц", 1),
}
UNIT_MAP = {1: 11, 4: 29, 5: 12, 8: 33}  # mln_rub, count, pct, months


def curl(url, out=None, retries=3):
    """Забрать через curl: попытки, со второй без проверки сертификата."""
    for attempt in range(1, retries + 1):
        cmd = ["curl", "-sL", "--max-time", "120", "-k" if attempt >= 2 else "-sS",
               "-A", "Mozilla/5.0 (research-wiki ingest watcher)", url]
        if out:
            cmd += ["-o", out, "-w", "%{http_code}"]
            r = subprocess.run(cmd, capture_output=True, text=True)
            if r.returncode == 0 and r.stdout.strip().startswith("2"):
                return True
        else:
            r = subprocess.run(cmd, capture_output=True)
            if r.returncode == 0 and r.stdout:
                for enc in ("utf-8", "cp1251", "latin-1"):
                    try:
                        return r.stdout.decode(enc)
                    except UnicodeDecodeError:
                        continue
        print(f"  попытка {attempt}: {url[:70]} не удалась", flush=True)
    return None


def datekey(v):
    if isinstance(v, datetime):
        return (v.year, v.month)
    if isinstance(v, str):
        m = re.match(r"(\S+)\s+(\d{4})", v.strip())
        if m and m.group(1) in MON:
            return (int(m.group(2)), MON[m.group(1)])
        if m and m.group(1).isdigit():
            return (int(m.group(2)), int(m.group(1)))
    return None


def num(x):
    if x is None:
        return None
    s = str(x).replace(" ", "").replace("\xa0", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def main():
    now = datetime.utcnow()
    today = now.strftime("%Y-%m-%d")
    body = curl(PAGE)
    if not body:
        print("ERROOR: страница ипотеки ЦБ недоступна")
        return 2
    urls = sorted({m.group(1) for m in
                   re.finditer(r'href=["\']([^"\']+\.xlsx)["\']', body, re.I)})
    urls = [u if u.startswith("http") else "https://www.cbr.ru" + u for u in urls]
    wanted = [u for u in urls if u.split("/")[-1][:5] in FORMS]
    print(f"найдено xlsx: {len(urls)}, нужных форм: {len(wanted)}")
    if not wanted:
        print("ERROR: ни одной нужной формы на странице — разметка сайта изменилась?")
        return 2

    conn = psycopg.connect(host=os.environ.get("PGHOST", "localhost"),
                           port=int(os.environ.get("PGPORT", "5432")),
                           dbname="research_wiki", user="wiki", autocommit=False)
    with conn.cursor() as cur:
        src_id = cur.execute("SELECT source_id FROM catalog.source WHERE source_code='cbr'").fetchone()[0]
        freq_m = cur.execute("SELECT frequency_id FROM catalog.frequency WHERE freq_code='M'").fetchone()[0]
        label = f"cbr-vfs-sync {today} (прямой канал, ingest-watcher)"
        rel = cur.execute("SELECT release_id FROM catalog.release WHERE label=%s", (label,)).fetchone()
        if rel:
            release_id = rel[0]
        else:
            release_id = cur.execute(
                """INSERT INTO catalog.release (source_id,label,published_at,status,notes)
                   VALUES (%s,%s,%s,'loaded','региональные формы 0420218: прямой канал cbr.ru')
                   RETURNING release_id""", (src_id, label, now.date())).fetchone()[0]
        conn.commit()
        rmap = {nm.strip().lower().replace("ё", "е"): rid for rid, nm in
                cur.execute("SELECT region_id, name_ru FROM catalog.region").fetchall()}
        midmap = {}
        for mc, (name, u2) in NAMES.items():
            r = cur.execute("SELECT metric_id FROM catalog.metric WHERE metric_code=%s", (mc,)).fetchone()
            if r:
                midmap[mc] = r[0]
            else:
                midmap[mc] = cur.execute(
                    """INSERT INTO catalog.metric
                       (metric_code,name_ru,unit_id,source_id,frequency_id,metric_type,status,is_active)
                       VALUES (%s,%s,%s,%s,%s,'primary','active',true) RETURNING metric_id""",
                    (mc, name, UNIT_MAP[u2], src_id, freq_m)).fetchone()[0]
                conn.commit()
        rows, rej_reg = [], set()
        for u in wanted:
            fname = u.split("/")[-1]
            prefix = fname[:5]
            path = f"/tmp/cbr_sync_{fname}"
            if not curl(u, out=path):
                print(f"ERROOR: файл не скачался {fname}")
                return 2
            with open(path, "rb") as f:
                raw = f.read()
            wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            for sn in wb.sheetnames:
                key = (prefix, sn.strip().lower())
                if key not in SHEET_MAP:
                    continue
                metric_code = SHEET_MAP[key]
                grid = list(wb[sn].iter_rows(values_only=True))
                hdr_i = next((i for i, r in enumerate(grid[:8])
                              if len([c for c in r[1:] if c is not None]) > 3), None)
                if hdr_i is None:
                    continue
                hdr = grid[hdr_i]
                n_file = 0
                for r_ in grid:
                    nm = str(r_[0]).strip() if r_[0] else ""
                    if not nm:
                        continue
                    kk = nm.lower().replace("ё", "е")
                    rid = rmap.get(kk) or (rmap.get(ALIAS[kk]) if kk in ALIAS else None)
                    if rid is None:
                        rej_reg.add(nm)
                        continue
                    for j in range(1, len(hdr)):
                        dk = datekey(hdr[j])
                        if dk is None:
                            continue
                        val = num(r_[j])
                        if val is None:
                            continue
                        rows.append((midmap[metric_code], rid, f"{dk[0]}-{dk[1]:02d}-01",
                                     freq_m, val, src_id, release_id))
                        n_file += 1
                print(f"  {fname}/{sn}: {n_file}", flush=True)
            wb.close()
            os.unlink(path)
        if not rows:
            print("ERROR: новых строк не собрано — структура листов изменилась?")
            return 2
        B = 4000
        inserted = 0
        for i in range(0, len(rows), B):
            cur.executemany(
                """INSERT INTO data.observation
                   (metric_id,region_id,period_start,frequency_id,value,
                    assessment_type,observation_status,source_id,release_id,sub_dimension)
                   VALUES (%s,%s,%s,%s,%s,'final','validated',%s,%s,'')
                   ON CONFLICT (metric_id,region_id,period_start,frequency_id,
                                assessment_type,sub_dimension,release_id) DO NOTHING""",
                rows[i:i + B])
            inserted += cur.rowcount
        conn.commit()
        print(f"релиз {release_id} ({label}): строк собрано {len(rows):,}, вставлено {inserted:,};"
              f" нераспознанные регионы: {sorted(rej_reg)[:6]} ({len(rej_reg)})")
    conn.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
#!/usr/bin/env python3
"""Прямой загрузчик: месячный бюллетень Росстата Prom_MM_YYYY.xlsx → канон.

Путь «сырьё → канон» вместо прежнего «парсер → SQLite-лендинг → гармонайзер 6 →
канон». Сырьё: raw/prom/Prom_MM_YYYY.xlsx (производство важнейших видов
продукции; листы Н-П_NN). Для каждого продукта берётся строка «Российская
Федерация без учёта …» с тремя графами: отчётный месяц, предыдущий месяц,
период с начала года.

Три метрики на продукт (имена совпадают с рубежом 6, чтобы попадать в те же
ключи):
    Промпродукция, месяц: __product__            (value_month)
    Промпродукция, пред. месяц: __product__      (value_prev_month)
    Промпродукция, накопленным итогом: __product__ (value_ytd)
Частота — M (месячная): рубеж 6 помечал эти ряды квартальной частотой (дефект,
исправлен 16.10.2026, см. attic/stage6_freq_before.csv.gz).

Запуск:  python3 scripts/direct_rosstat_prom.py [файл.xlsx ...]
         SIMULATE=1 python3 scripts/direct_rosstat_prom.py   # без записи
"""
import glob
import os
import re
import sys

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from direct_core import Writer, norm  # noqa: E402

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CODE_RE = re.compile(r'^\d{2}(\.\d+)+$')
RF_PREFIX = 'Российская Федерация'
SOURCE = 'rosstat'
FREQ = 'M'
SKIP_SHEETS = ('Содержание', 'Лист1')


def num(v):
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        try:
            return float(v.strip().replace(' ', '').replace(',', '.'))
        except ValueError:
            return None
    return None


def parse_sheet(ws):
    """(продукт, единица, отчётный месяц, пред. месяц, накопленный) для строк РФ."""
    out = []
    product, unit = None, None
    for row in ws.iter_rows(values_only=True):
        a = row[0] if len(row) > 0 else None
        b = row[1] if len(row) > 1 else None
        if isinstance(a, str) and isinstance(b, str) and CODE_RE.match(b.strip()) and len(a.strip()) > 3:
            product, unit = a.strip(), None
            continue
        if product and unit is None and isinstance(a, str) and a.strip():
            bs = str(b).strip() if b is not None else ''
            if bs.isdigit() and len(bs) <= 6:      # строка единицы измерения с кодом ОКЕИ
                unit = a.strip()
                continue
        if isinstance(a, str) and a.strip().startswith(RF_PREFIX) and product:
            vals = [num(row[i]) if len(row) > i else None for i in (2, 3, 4)]
            if any(v is not None for v in vals):
                out.append((product, unit, vals[0], vals[1], vals[2]))
            product, unit = None, None             # один РФ-ряд на блок продукта
    return out


def load_file(path, w):
    m = re.search(r'Prom_(\d{2})_(\d{4})', os.path.basename(path))
    if not m:
        print('ПРОПУСК: не разобрать месяц/год из имени', os.path.basename(path))
        return [], None
    month, year = int(m.group(1)), int(m.group(2))
    label = 'rosstat_Prom_%02d_%04d' % (month, year)
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    rows = []
    for sheet in wb.sheetnames:
        if sheet in SKIP_SHEETS:
            continue
        for product, unit, vm, vp, vytd in parse_sheet(wb[sheet]):
            sub = 'product: %s' % product
            if unit and str(unit).strip() not in ('', 'None', 'ND'):
                sub += ' | ед.: %s' % str(unit)[:40]
            for tmpl, val in (('Промпродукция, месяц: %s', vm),
                              ('Промпродукция, пред. месяц: %s', vp),
                              ('Промпродукция, накопленным итогом: %s', vytd)):
                rows.append(dict(metric_name=tmpl % product, value=val, ps=(year, month),
                                 region=w.RU, sub=sub, freq=FREQ, source=SOURCE,
                                 release=label, unit='unknown'))
    wb.close()
    return rows, label


def main():
    args = [a for a in sys.argv[1:] if not a.startswith('-')]
    files = args or sorted(glob.glob(os.path.join(REPO, 'raw', 'prom', 'Prom_*.xlsx')))
    if not files:
        sys.exit('не найдено raw/prom/Prom_*.xlsx')
    w = Writer(tag='direct_prom', simulate=os.environ.get('SIMULATE') == '1')
    total_prep = total_ins = 0
    for f in files:
        if not os.path.exists(f):
            print('ПРОПУСК: нет файла', f)
            continue
        rows, label = load_file(f, w)
        if not label:
            continue
        prep, ins, conf, nm = w.ingest(rows, group=label)
        print('%-26s подготовлено %6d | вставлено %6d | конфликтов %6d | новых метрик %d'
              % (os.path.basename(f), prep, ins, conf, nm), flush=True)
        total_prep += prep
        total_ins += ins
    print('ИТОГО подготовлено %d, вставлено %d' % (total_prep, total_ins))
    w.close()


if __name__ == '__main__':
    main()

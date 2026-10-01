#!/usr/bin/env python3
"""Добор рядов первичного рынка ЦБ: ИЖК под залог прав требования по ДДУ
(файлы 02_15/02_16/02_17) и ИЖК на цели ИЖС (02_42/02_43/02_44/02_45).

Этап 0 дизайна queries/dizayn-prognoz-spros-pervichka-reforma-2026.md.
Задача канбана t_d3dc28ff.

Что делает:
  1. скачивает 7 файлов ЦБ в raw/cbr/vfs/ (кеш, sha256);
  2. разбирает листы «в рублях» / «ставка в рублях» в длинный формат;
  3. создаёт и наполняет public.cbr_ihk_ddu и public.cbr_ihk_izhs
     (схема как у public.cbr_housing_loans: level/fed_okrug/region/period/metric/value);
  4. регистрирует в v2 (dataset, metric, metric_dataset_map, load_run, load_check,
     freshness) и в core.metric + derived.catalog_facets (для видимости в
     v_catalog_active — ловушка из migrate_v2_core_catalog.sql);
  5. прогоняет приёмку: число субъектов, глубина, сумма субъектов против строки
     «РОССИЙСКАЯ ФЕДЕРАЦИЯ», разложение автономных округов, контрольные точки.

Идемпотентно: повторный запуск не плодит строк (DELETE + INSERT по таблице,
ON CONFLICT DO NOTHING в каталоге).

Запуск:
    ~/.hermes/hermes-agent/venv/bin/python3 scripts/load_cbr_primary_market.py
"""
import hashlib
import json
import os
import re
import sys
import urllib.request
import ssl
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
import db_tunnel as dbt

ROOT = Path(__file__).resolve().parent.parent
VFS_DIR = ROOT / "raw" / "cbr" / "vfs"
BASE_URL = "https://www.cbr.ru/vfs/statistics/BankSector/Mortgage/%s.xlsx"

MON = {"Январь": 1, "Февраль": 2, "Март": 3, "Апрель": 4, "Май": 5, "Июнь": 6,
       "Июль": 7, "Август": 8, "Сентябрь": 9, "Октябрь": 10, "Ноябрь": 11,
       "Декабрь": 12}

# (таблица, файл, лист, метрика)
FILES = [
    ("cbr_ihk_ddu",  "02_15_Quantity_scpa_mortgage",   "в рублях",       "count"),
    ("cbr_ihk_ddu",  "02_16_New_loans_scpa_mortgage",  "в рублях",       "volume"),
    ("cbr_ihk_ddu",  "02_17_Rates_scpa_mortgage",      "ставка в рублях", "rate"),
    ("cbr_ihk_izhs", "02_42_Quantity_mortgage_ihc",    "в рублях",       "count"),
    ("cbr_ihk_izhs", "02_43_New_loans_mortgage_ihc",   "в рублях",       "volume"),
    ("cbr_ihk_izhs", "02_44_Share_mortgage_ihc",       "в рублях",       "share"),
    ("cbr_ihk_izhs", "02_45_Rates_mortgage_ihc",       "ставка в рублях", "rate"),
]

# Регистрация в v2 / core. (v2_id, core_id, metric_code, name_ru, unit_v2, unit_core,
#                          unit_canonical, table, metric, is_agg)
METRICS = [
    (58065, 71947, "cbr_ihk_ddu_count_reg_m",
     "ИЖК под залог прав требования по ДДУ, количество (рег.)",
     4, 29, "count", "cbr_ihk_ddu", "count", False),
    (58066, 71948, "cbr_ihk_ddu_volume_reg_m",
     "ИЖК под залог прав требования по ДДУ, объём (рег.)",
     1, 11, "mln_rub", "cbr_ihk_ddu", "volume", False),
    (58067, 71949, "cbr_ihk_ddu_rate_reg_m",
     "ИЖК под залог прав требования по ДДУ, ставка (рег.)",
     5, 12, "pct", "cbr_ihk_ddu", "rate", False),
    (58068, 71950, "cbr_ihk_izhs_count_reg_m",
     "ИЖК на приобретение и создание объектов ИЖС, количество (рег.)",
     4, 29, "count", "cbr_ihk_izhs", "count", False),
    (58069, 71951, "cbr_ihk_izhs_volume_reg_m",
     "ИЖК на приобретение и создание объектов ИЖС, объём (рег.)",
     1, 11, "mln_rub", "cbr_ihk_izhs", "volume", False),
    (58070, 71952, "cbr_ihk_izhs_share_reg_m",
     "Доля ИЖК на цели ИЖС в общем объёме ИЖК (рег.)",
     5, 12, "pct", "cbr_ihk_izhs", "share", False),
    (58071, 71953, "cbr_ihk_izhs_rate_reg_m",
     "ИЖК на приобретение и создание объектов ИЖС, ставка (рег.)",
     5, 12, "pct", "cbr_ihk_izhs", "rate", False),
]

RUN_TAG = "cbr-primary-market-20261001"


# ── разбор ────────────────────────────────────────────────────────────────
def month_key(v):
    if v is None:
        return None
    s = str(v).strip()
    m = re.match(r"^(\S+)\s+(\d{4})$", s)
    if m and m.group(1) in MON:
        return "%s-%02d" % (m.group(2), MON[m.group(1)])
    m = re.match(r"^(\d{2})\.(\d{2})\.(\d{4})$", s)
    if m:
        return "%s-%s" % (m.group(3), m.group(2))
    m = re.match(r"^(\d{4})-(\d{2})", s)
    if m:
        return "%s-%s" % (m.group(1), m.group(2))
    return None


def num(x):
    if x is None:
        return None
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).replace(" ", "").replace("\xa0", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def _norm(s):
    s = str(s).lower().replace("\n", " ")
    s = re.sub(r"\s+", " ", s).strip()
    return s.replace("h", "н")


def _parse_bulletin(path, sheet):
    """Разбор листа бюллетеня ЦБ (месяцы по столбцам) в dict {(норм.регион, период): value}."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        wb.close()
        return {}
    grid = list(wb[sheet].iter_rows(values_only=True))
    wb.close()
    hdr_i = None
    for i, r in enumerate(grid[:8]):
        if sum(1 for c in r[1:] if month_key(c)) >= 3:
            hdr_i = i
            break
    if hdr_i is None:
        return {}
    hdr = grid[hdr_i]
    cols = [(j, month_key(hdr[j])) for j in range(1, len(hdr)) if month_key(hdr[j])]
    out = {}
    for r in grid[hdr_i + 1:]:
        nm = r[0]
        if nm is None:
            continue
        nm = _norm(nm)
        for j, per in cols:
            v = num(r[j]) if j < len(r) else None
            if v is not None:
                out[(nm, per)] = v
    return out


def classify(label):
    """Уровень строки. Возвращает (level, region_name)."""
    lab = label.strip()
    up = lab.upper()
    if up == "РОССИЙСКАЯ ФЕДЕРАЦИЯ":
        return "RF", lab
    if up.endswith("ФЕДЕРАЛЬНЫЙ ОКРУГ"):
        return "federal_district", lab
    if lab.lower().startswith("в том числе"):
        # автономные округа внутри материнской области (субъекты РФ).
        # ВАЖНО: «Ямало-Ненецкий» содержит подстроку «ненецкий» — порядок проверок.
        low = lab.lower()
        if "ханты" in low:
            return "okrug", "ХМАО-Югра"
        if "ямало" in low:
            return "okrug", "ЯНАО"
        if "ненецкий" in low:
            return "okrug", "НАО"
        return "okrug", lab
    if "без данных" in lab.lower():
        if lab.lower().startswith("архангельская"):
            return "nested", "Архангельская область (без НАО)"
        if lab.lower().startswith("тюменская"):
            return "nested", "Тюменская область (без Югры и ЯНАО)"
        return "nested", lab
    return "region", lab


def parse_sheet(path, sheet, metric):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    grid = list(wb[sheet].iter_rows(values_only=True))
    wb.close()
    hdr_i = None
    for i, r in enumerate(grid[:8]):
        if sum(1 for c in r[1:] if month_key(c)) >= 5:
            hdr_i = i
            break
    if hdr_i is None:
        raise RuntimeError("no header row in %s / %s" % (path, sheet))
    hdr = grid[hdr_i]
    cols = [(j, month_key(hdr[j])) for j in range(1, len(hdr)) if month_key(hdr[j])]
    out = []
    cur_fo = None
    for r in grid[hdr_i + 1:]:
        nm = r[0]
        if nm is None:
            continue
        nm = str(nm).replace("\n", " ").strip()
        nm = re.sub(r"\s+", " ", nm)
        if not nm:
            continue
        level, region = classify(nm)
        if level == "federal_district":
            cur_fo = region
            fo = region
        elif level == "RF":
            fo = "Российская Федерация"
        else:
            fo = cur_fo
        for j, per in cols:
            v = num(r[j]) if j < len(r) else None
            if v is None:
                continue
            out.append((level, fo, region, per, metric, v))
    return out


def download(fname):
    dst = VFS_DIR / (fname + ".xlsx")
    url = BASE_URL % fname
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    ctx = ssl.create_default_context()
    with urllib.request.urlopen(req, timeout=90, context=ctx) as r:
        data = r.read()
    dst.write_bytes(data)
    return dst, hashlib.sha256(data).hexdigest()


# ── загрузка ──────────────────────────────────────────────────────────────
DDL = """
CREATE TABLE IF NOT EXISTS public.{tbl} (
    level      text NOT NULL,
    fed_okrug  text,
    region     text NOT NULL,
    period     text NOT NULL,
    metric     text NOT NULL,
    value      double precision
);
CREATE INDEX IF NOT EXISTS {tbl}_region_period_idx ON public.{tbl} (region, period);
CREATE INDEX IF NOT EXISTS {tbl}_metric_idx ON public.{tbl} (metric);
"""


def main():
    dbt.available(verbose=True)
    VFS_DIR.mkdir(parents=True, exist_ok=True)

    parsed = {}      # table -> list of rows
    file_meta = []   # (table, metric, fname, sha, nrows, months)
    for tbl, fname, sheet, metric in FILES:
        path, sha = download(fname)
        rows = parse_sheet(path, sheet, metric)
        months = sorted({r[3] for r in rows})
        parsed.setdefault(tbl, []).extend(rows)
        file_meta.append((tbl, metric, fname, sha, len(rows), months[0], months[-1], len(months)))
        print("%-14s %-34s %-16s rows=%-6d months=%d (%s..%s)" %
              (tbl, fname, metric, len(rows), len(months), months[0], months[-1]))

    # таблицы
    with dbt.connect() as c:
        with c.cursor() as cur:
            for tbl in parsed:
                cur.execute(DDL.format(tbl=tbl))
                cur.execute("TRUNCATE public.%s" % tbl)
                cur.execute("GRANT SELECT ON public.%s TO viewer" % tbl)
            for tbl, rows in parsed.items():
                cur.executemany(
                    "INSERT INTO public.%s (level, fed_okrug, region, period, metric, value) "
                    "VALUES (%%s,%%s,%%s,%%s,%%s,%%s)" % tbl, rows)
        c.commit()

    # ── v2 регистрация ────────────────────────────────────────────────────
    with dbt.connect() as c:
        with c.cursor() as cur:
            cur.execute("SELECT dataset_id FROM v2.dataset WHERE dataset_code='cbr-primary-market'")
            row = cur.fetchone()
            if row:
                ds_id = row[0]
                cur.execute("UPDATE v2.dataset SET name_ru=%s, is_active=true WHERE dataset_id=%s",
                            ("ЦБ РФ — ИЖК под залог прав требования по ДДУ и ИЖК на цели ИЖС (региональный разрез)", ds_id))
            else:
                cur.execute("SELECT coalesce(max(dataset_id),0)+1 FROM v2.dataset")
                ds_id = cur.fetchone()[0]
                cur.execute("""INSERT INTO v2.dataset (dataset_id, source_id, dataset_code, name_ru, release_url, is_active)
                               VALUES (%s,
                                       (SELECT source_id FROM v2.source WHERE source_code='cbr'),
                                       'cbr-primary-market',
                                       'ЦБ РФ — ИЖК под залог прав требования по ДДУ и ИЖК на цели ИЖС (региональный разрез)',
                                       'https://www.cbr.ru/statistics/bank_sector/mortgage/', true)""",
                            (ds_id,))

            for v2_id, core_id, code, name, u2, ucore, ucanon, tbl, metric, is_agg in METRICS:
                cur.execute("""INSERT INTO v2.metric
                    (metric_id, metric_code, name_ru, definition_ru, domain_id, unit_id,
                     frequency_id, is_aggregate, priority)
                    VALUES (%s,%s,%s,%s,4,%s,2,%s,'primary')
                    ON CONFLICT (metric_id) DO NOTHING""",
                    (v2_id, code, name,
                     "Источник: ЦБ РФ, файлы 02_15-02_17 (ДДУ) / 02_42-02_45 (ИЖС). "
                     "Таблица public.%s, метрика '%s'." % (tbl, metric), u2, is_agg))
                cur.execute("""INSERT INTO v2.metric_dataset_map
                    (metric_id, dataset_id, native_code, scale_factor)
                    VALUES (%s,%s,%s,1)
                    ON CONFLICT (metric_id, dataset_id) DO NOTHING""",
                    (v2_id, ds_id, "%s:%s" % (tbl, metric)))

            # load_run по каждому файлу
            cur.execute("SELECT coalesce(max(run_id),0) FROM v2.load_run")
            next_run = cur.fetchone()[0] + 1
            cur.execute("SELECT run_id, source_file FROM v2.load_run WHERE dataset_id=%s", (ds_id,))
            existing = {r[1]: r[0] for r in cur.fetchall()}
            run_ids = []
            for tbl, metric, fname, sha, nrows, mn, mx, ndep in file_meta:
                sf = "%s.xlsx" % fname
                if sf in existing:
                    run_ids.append(existing[sf])
                    continue
                cur.execute("""INSERT INTO v2.load_run
                    (run_id, dataset_id, source_file, file_sha256, source_published_at, parser_version,
                     status, rows_parsed, rows_inserted, notes)
                    VALUES (%s,%s,%s,%s,%s,'load_cbr_primary_market@v1','ok',%s,%s,%s)
                    RETURNING run_id""",
                    (next_run, ds_id, sf, sha, (mx + "-01") if re.match(r"\d{4}-\d{2}$", mx) else None,
                     nrows, nrows, "таблица public.%s, метрика %s, глубина %s..%s" % (tbl, metric, mn, mx)))
                run_ids.append(cur.fetchone()[0])
                next_run += 1
        c.commit()

    print("\nv2 dataset_id =", ds_id, "load_run ids =", run_ids)
    return ds_id, run_ids, parsed, file_meta


CORE_TAGS = ["v2", "housing", "mortgage", "cbr", "primary"]


def register_core(ds_id, file_meta):
    """core.metric + catalog_facets, чтобы ряды были видны в v_catalog_active."""
    depth = {}
    for tbl, metric, fname, sha, nrows, mn, mx, ndep in file_meta:
        depth[(tbl, metric)] = (ndep, mx)
    with dbt.connect() as c:
        with c.cursor() as cur:
            for v2_id, core_id, code, name, u2, ucore, ucanon, tbl, metric, is_agg in METRICS:
                cur.execute("""INSERT INTO core.metric
                    (metric_id, metric_code, name_ru, short_name_ru, description,
                     unit_id, frequency_id, metric_type, is_derived, status, tags)
                    OVERRIDING SYSTEM VALUE
                    VALUES (%s,%s,%s,%s,%s,%s,5,'primary',false,'active',%s)
                    ON CONFLICT (metric_id) DO NOTHING""",
                    (core_id, code, name, name[:60],
                     "ЦБ РФ, файлы 02_15-02_17 (ДДУ) / 02_42-02_45 (ИЖС); "
                     "таблица public.%s, метрика '%s'." % (tbl, metric),
                     ucore, CORE_TAGS))
                ndep, mx = depth[(tbl, metric)]
                cur.execute("""INSERT INTO derived.catalog_facets
                    (metric_id, measure, unit_canonical, seasonally_adjusted, domain,
                     freshness, last_obs_at, depth_obs, is_aggregate, run_id)
                    VALUES (%s,'level',%s,false,'credit','fresh',%s,%s,%s,%s)
                    ON CONFLICT (metric_id) DO NOTHING""",
                    (core_id, ucanon, mx + "-01", ndep, is_agg, RUN_TAG))
                cur.execute("""INSERT INTO v2.freshness
                    (metric_id, last_obs_at, obs_count, freshness, expected_next)
                    VALUES (%s,%s,%s,'fresh',%s)
                    ON CONFLICT (metric_id) DO UPDATE
                      SET last_obs_at=EXCLUDED.last_obs_at, obs_count=EXCLUDED.obs_count,
                          freshness=EXCLUDED.freshness, updated_at=now()""",
                    (v2_id, mx + "-01", ndep, None))
            cur.execute("SELECT setval(pg_get_serial_sequence('core.metric','metric_id'), "
                        "(SELECT max(metric_id) FROM core.metric))")
        c.commit()
    print("core.metric + catalog_facets + v2.freshness зарегистрированы")


def acceptance(run_ids):
    """Приёмка. Возвращает dict с результатами; пишет v2.load_check."""
    checks = []

    def add(kind, passed, detail):
        checks.append((kind, bool(passed), detail))
        print("  [%s] %-42s %s" % ("OK " if passed else "FAIL", kind, detail))

    print("\n=== ПРИЁМКА ===")
    for tbl in ("cbr_ihk_ddu", "cbr_ihk_izhs"):
        n_subj = dbt.query("SELECT count(DISTINCT region) FROM public.%s "
                           "WHERE level IN ('region','okrug')" % tbl)[0][0]
        n_reg = dbt.query("SELECT count(DISTINCT region) FROM public.%s WHERE level='region'" % tbl)[0][0]
        add("%s:субъектов>=85" % tbl, n_subj >= 85, "субъектов=%d (region=%d)" % (n_subj, n_reg))
        nmon = dbt.query("SELECT count(DISTINCT period) FROM public.%s" % tbl)[0][0]
        exp = 92 if tbl == "cbr_ihk_ddu" else 32
        add("%s:глубина=%d" % (tbl, exp), nmon == exp, "месяцев=%d" % nmon)
        # сумма субъектов против строки РФ, по каждой метрике
        for (metric,) in dbt.query("SELECT DISTINCT metric FROM public.%s ORDER BY 1" % tbl):
            if metric in ("share", "rate"):   # доля и средневзвешенная ставка не суммируются
                continue
            bad = dbt.query("""
                WITH rf AS (SELECT period, value FROM public.{t} WHERE level='RF' AND metric=%s),
                     sm AS (SELECT period, sum(value) v FROM public.{t} WHERE level='region' AND metric=%s GROUP BY period)
                SELECT count(*) FROM rf JOIN sm USING(period)
                WHERE abs(rf.value - sm.v) > greatest(0.005*abs(rf.value), 1.5)
            """.format(t=tbl), (metric, metric))[0][0]
            add("%s:%s сумма(region)=РФ" % (tbl, metric), bad == 0,
                "расхождений сверх допуска: %d из %d мес" % (bad, nmon))
        # разложение автономных округов
        for whole, parts, name in [
                ("Архангельская область",
                 ["Архангельская область (без НАО)", "НАО"], "Архангельская"),
                ("Тюменская область",
                 ["Тюменская область (без Югры и ЯНАО)", "ХМАО-Югра", "ЯНАО"], "Тюменская")]:
            bad = dbt.query("""
                WITH w AS (SELECT period, value FROM public.{t} WHERE region=%s AND metric='count'),
                     p AS (SELECT period, sum(value) v FROM public.{t}
                           WHERE region = ANY(%s) AND metric='count' GROUP BY period)
                SELECT count(*) FROM w JOIN p USING(period) WHERE abs(w.value - p.v) > 1.5
            """.format(t=tbl), (whole, parts))[0][0]
            add("%s:разложение %s" % (tbl, name), bad == 0, "расхождений: %d" % bad)

    # контрольные точки (значения строки РФ)
    for tbl, metric, period, exp, tol in [
            ("cbr_ihk_ddu", "count", "2019-01", 22845.0, 0.5),
            ("cbr_ihk_ddu", "count", "2026-08", 28279.0, 0.5),
            ("cbr_ihk_ddu", "volume", "2026-08", 161511.0, 0.5),
            ("cbr_ihk_ddu", "rate", "2019-01", 9.65, 0.01),
            ("cbr_ihk_izhs", "count", "2024-01", 22416.0, 0.5),
            ("cbr_ihk_izhs", "count", "2026-08", 13111.0, 0.5),
            ("cbr_ihk_izhs", "rate", "2024-01", 7.96, 0.01)]:
        got = dbt.query("SELECT value FROM public.%s WHERE level='RF' AND metric=%%s AND period=%%s"
                        % tbl, (metric, period))
        got = got[0][0] if got else None
        add("контроль %s/%s/%s" % (tbl, metric, period),
            got is not None and abs(got - exp) <= tol, "файл=%.4g ожидание=%.4g" % (got or -1, exp))

    # записать load_check (идемпотентно: снять прежние проверки этих прогонов)
    with dbt.connect() as c:
        with c.cursor() as cur:
            cur.execute("DELETE FROM v2.load_check WHERE run_id = ANY(%s) AND checked_by='agent'", (list(run_ids),))
            cur.execute("SELECT coalesce(max(check_id),0) FROM v2.load_check")
            nxt = cur.fetchone()[0] + 1
            for run_id in run_ids:
                for kind, passed, detail in checks:
                    cur.execute("""INSERT INTO v2.load_check (check_id, run_id, check_kind, check_detail, passed, checked_by)
                                   VALUES (%s,%s,%s,%s,%s,'agent')""",
                                (nxt, run_id, kind, json.dumps({"detail": detail}, ensure_ascii=False), passed))
                    nxt += 1
        c.commit()
    n_pass = sum(1 for _, p, _ in checks if p)
    print("\nитог: %d/%d проверок пройдено" % (n_pass, len(checks)))
    return {"passed": n_pass, "total": len(checks)}


def crosscheck_bulletin(run_ids):
    """Внешняя сверка: контрольные точки против бюллетеня ЦБ
    «Ипотечное жилищное кредитование» (листы Т_17/Т_20/Т_27 — независимая
    публикация ЦБ по тем же показателям). Только для месяцев, попадающих в
    окно бюллетеня (13 мес)."""
    bul_dir = ROOT / "raw" / "cbr" / "bulletin"
    files = sorted(bul_dir.glob("mortgage_lending_market_*.xlsx")) if bul_dir.exists() else []
    if not files:
        print("\n=== ВНЕШНЯЯ СВЕРКА: файлы бюллетеня не найдены, пропуск ===")
        return None
    mine = {}
    for metric in ("count", "volume", "rate"):
        for reg, per, val in dbt.query("SELECT region, period, value FROM public.cbr_ihk_ddu WHERE metric=%s", (metric,)):
            mine[(metric, per, _norm(reg))] = val
    total = ok = 0
    for bfile in files:
        for sheet, metric in [("Т_17", "count"), ("T_17", "count"), ("Т_20", "volume"),
                              ("T_20", "volume"), ("Т_27", "rate"), ("T_27", "rate")]:
            data = _parse_bulletin(bfile, sheet)
            for (nm, per), v in data.items():
                key = (metric, per, nm)
                if key in mine:
                    total += 1
                    if abs(mine[key] - v) <= max(0.01 * abs(v), 0.05):
                        ok += 1
    passed = total > 0 and ok >= 0.99 * total
    print("\n=== ВНЕШНЯЯ СВЕРКА (бюллетень ЦБ Т_17/Т_20/Т_27) ===")
    print("  пар сверено: %d, совпало: %d (%.2f%%)" % (total, ok, 100.0 * ok / total if total else 0))
    with dbt.connect() as c:
        with c.cursor() as cur:
            cur.execute("DELETE FROM v2.load_check WHERE run_id = ANY(%s) AND check_kind LIKE 'внешняя сверка%%'", (list(run_ids),))
            cur.execute("SELECT coalesce(max(check_id),0) FROM v2.load_check")
            nxt = cur.fetchone()[0] + 1
            for run_id in run_ids:
                cur.execute("""INSERT INTO v2.load_check (check_id, run_id, check_kind, check_detail, passed, checked_by)
                               VALUES (%s,%s,'внешняя сверка бюллетень ЦБ',%s,%s,'agent')""",
                            (nxt, run_id, json.dumps({"pairs": total, "ok": ok}, ensure_ascii=False), passed))
                nxt += 1
        c.commit()
    return {"pairs": total, "ok": ok, "passed": passed}


if __name__ == "__main__":
    ds_id, run_ids, parsed, file_meta = main()
    register_core(ds_id, file_meta)
    acceptance(run_ids)
    crosscheck_bulletin(run_ids)
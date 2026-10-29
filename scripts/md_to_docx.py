#!/usr/bin/env python3
"""Универсальный конвертер markdown -> docx (подмножество разметки обзоров).

Использование: md_to_docx.py <input.md> <output.docx>
Поддерживает: YAML-frontmatter (пропускается), # / ## / ###, списки '- ',
нумерованные '1. ', **жирный**, курсив *...*, инлайн-код `...`.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, Mm


def add_runs(paragraph, text):
    parts = re.split(r'(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith('`') and part.endswith('`') and len(part) > 2:
            r = paragraph.add_run(part[1:-1])
            r.font.name = 'Consolas'
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def main():
    inp, out = Path(sys.argv[1]), Path(sys.argv[2])
    lines = inp.read_text(encoding="utf-8").splitlines()

    # пропустить YAML-frontmatter
    if lines and lines[0].strip() == '---':
        for k in range(1, len(lines)):
            if lines[k].strip() == '---':
                lines = lines[k + 1:]
                break

    doc = Document()
    section = doc.sections[0]
    section.page_height, section.page_width = Mm(297), Mm(210)
    section.top_margin = section.bottom_margin = Mm(20)
    section.left_margin, section.right_margin = Mm(25), Mm(15)

    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.15

    lines_all = lines
    idx = 0
    while idx < len(lines_all):
        stripped = lines_all[idx].strip()
        # таблица: подряд идущие строки, начинающиеся с '|'
        if stripped.startswith('|'):
            block = []
            while idx < len(lines_all) and lines_all[idx].strip().startswith('|'):
                block.append(lines_all[idx].strip())
                idx += 1
            rows = []
            for r in block:
                cells = [c.strip() for c in r.strip('|').split('|')]
                if all(set(c) <= set('-: ') and c for c in cells):
                    continue
                rows.append(cells)
            if rows:
                ncol = max(len(r) for r in rows)
                table = doc.add_table(rows=0, cols=ncol)
                table.style = 'Table Grid'
                for ri, r in enumerate(rows):
                    cells = table.add_row().cells
                    for ci in range(ncol):
                        txt = r[ci] if ci < len(r) else ''
                        cells[ci].text = ''
                        p = cells[ci].paragraphs[0]
                        add_runs(p, txt)
                        if ri == 0:
                            for run in p.runs:
                                run.bold = True
            continue

        line = lines_all[idx].rstrip()
        idx += 1
        stripped = line.strip()
        if not stripped or stripped == '---':
            continue
        if stripped.startswith('#### '):
            doc.add_heading(stripped[5:].strip(), level=4)
            continue
        if stripped.startswith('### '):
            doc.add_heading(stripped[4:].strip(), level=3)
            continue
        if stripped.startswith('## '):
            doc.add_heading(stripped[3:].strip(), level=2)
            continue
        if stripped.startswith('# '):
            doc.add_heading(stripped[2:].strip(), level=1)
            continue
        if stripped.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            add_runs(p, stripped[2:].strip())
            continue
        m = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if m:
            p = doc.add_paragraph(style='List Number')
            add_runs(p, m.group(2))
            continue
        p = doc.add_paragraph()
        add_runs(p, stripped)

    doc.save(str(out))
    print(f"OK -> {out} ({out.stat().st_size} bytes)")


if __name__ == "__main__":
    main()

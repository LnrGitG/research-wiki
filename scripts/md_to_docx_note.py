#!/usr/bin/env python3
"""Convert the CBR supervision note markdown to docx.

Reads queries/cbr-supervision-note-nolinks-20260929.md, produces
queries/cbr-supervision-note-nolinks-20260930.docx.

Markdown subset supported:
  # / ## / ###  -> Heading 1/2/3
  - item        -> bullet list
  blank line    -> paragraph break
  **bold**      -> bold run
  plain text    -> Normal paragraph
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH

INP = Path(__file__).parent / "../queries/cbr-supervision-note-nolinks-20260929.md"
OUT = Path(__file__).parent / "../queries/cbr-supervision-note-nolinks-20260930.docx"


def add_runs(paragraph, text):
    """Split text on **bold** markers and add runs."""
    parts = re.split(r'(\*\*[^*]+\*\*)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            r = paragraph.add_run(part[2:-2])
            r.bold = True
        else:
            paragraph.add_run(part)


def main():
    lines = INP.read_text(encoding="utf-8").splitlines()

    doc = Document()

    # Page setup
    section = doc.sections[0]
    section.page_height = Mm(297)
    section.page_width = Mm(210)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(25)
    section.right_margin = Mm(15)

    # Base font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.15

    i = 0
    while i < len(lines):
        line = lines[i]

        # Heading
        if line.startswith('### '):
            h = doc.add_heading(line[4:].strip(), level=3)
            i += 1
            continue
        if line.startswith('## '):
            h = doc.add_heading(line[3:].strip(), level=2)
            i += 1
            continue
        if line.startswith('# '):
            h = doc.add_heading(line[2:].strip(), level=1)
            i += 1
            continue

        # Bullet list
        if line.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            add_runs(p, line[2:].strip())
            i += 1
            continue

        # Blank line
        if not line.strip():
            i += 1
            continue

        # Normal paragraph — collect consecutive non-empty, non-special lines
        para_lines = [line.strip()]
        j = i + 1
        while j < len(lines):
            nxt = lines[j]
            if (not nxt.strip()
                    or nxt.startswith('#')
                    or nxt.startswith('- ')):
                break
            para_lines.append(nxt.strip())
            j += 1
        text = ' '.join(para_lines)
        p = doc.add_paragraph()
        add_runs(p, text)
        i = j

    doc.save(str(OUT))
    print(f"OK -> {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
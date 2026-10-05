#!/usr/bin/env python3
"""Канон-перевод: ФНС 1-НОМ квартальный срез (reg.xls/xlsx с сайта ФНС) → catalog + data.observation.

Этап 3 единый приёмный контур. Лист 1010 («всего» по видам налогов),
графы 1..40 → метрики fns_1nom_1..40. Уровни: РФ (region 1), ФО
(aggregation='source_aggregate', вставляются — решение владельца
06.10.2026), субъекты. Единица 15 = тыс. руб., частота 3 = годовая
(нарастающий итог на дату среза), период = дата среза (2026-04-01 /
2026-07-01). Строки «в том числе:» пропускаются. Графы 41–43 (пени)
вне канона — пропускаются.

Использование:
  python3 parse_fns_1nom_canon.py <xls_файл> <label> <YYYY-MM-DD>
  python3 parse_fns_1nom_canon.py /tmp/1nom010726reg.xls fns_1nom_20260701 2026-07-01

Точка возврата: DELETE FROM data.observation WHERE release_id=<id>.
Идемпотентность: релиз ищется по label; наблюдения ON CONFLICT DO
NOTHING по уникальному индексу data.observation.
"""
import re
import sys

import requests

BASE = ("/html/sites/www.new.nalog.ru/files/related_activities/"
        "statistics_and_analytics/forms/1nom/")
URL = "https://www.nalog.gov.ru" + BASE

FREQ_YEAR = 3
UNIT_THS_RUB = 15
SOURCE_CODE = "fns"
COLS = list(range(1, 41))
SKIP_LEVELS = ("в том числе:", "справочно", "из них")

ALIAS = {
    'республика марий-эл': 'республика марий эл',
    'республика северная осетия-алания': 'республика северная осетия — алания',
    'республика северная осетия - алания': 'республика северная осетия — алания',
    'чувашская республика-чувашия': 'чувашская республика',
    'республика татарстан (татарстан)': 'республика татарстан',
    'республика адыгея (адыгея)': 'республика адыгея',
    'кемеровская область': 'кемеровская область — кузбасс',
    'кемеровская область - кузбасс': 'кемеровская область — кузбасс',
    'ханты-мансийский автономный округ - югра': 'ханты-мансийский автономный округ — югра',
    'ханты-мансийский ао - югра': 'ханты-мансийский автономный округ — югра',
    'ненецкий ао': 'ненецкий автономный округ',
    'ямало-ненецкий ао': 'ямало-ненецкий автономный округ',
    'чукотский ао': 'чукотский автономный округ',
    'архангельская область': 'архангельская область (без нао)',
    'тюменская область': 'тюменская область (без хмао и янао)',
    'г. москва': 'москва',
    'г.москва': 'москва',
    'г. санкт-петербург': 'санкт-петербург',
    'г.санкт-петербург': 'санкт-петербург',
    'г. севастополь': 'севастополь',
    'г.севастополь': 'севастополь',
}


def norm(name):
    s = re.sub(r'\s+', ' ', str(name or '')).strip().lower()
    return ALIAS.get(s, s)


def read_sheet(path):
    """Лист 1010 → {(region_raw, col): value}. Формат xls или xlsx по расширению."""
    if path.lower().endswith(('.xlsx', '.xlsm')):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb['1010']
        rows = list(ws.iter_rows(values_only=True))
    else:
        import xlrd
        wb = xlrd.open_workbook(path)
        s = wb.sheet_by_name('1010')
        rows = [[s.cell_value(i, j) for j in range(s.ncols)]
                for i in range(s.nrows)]
    cells = {}
    for rr in rows[8:]:
        name = rr[0]
        if name is None:
            continue
        nm = str(name).strip()
        if not nm or any(nm.lower().startswith(p) for p in SKIP_LEVELS):
            continue
        vals = rr[3:46]
        for gi, v in enumerate(vals, start=1):
            if gi not in COLS:
                continue
            if v is None or str(v).strip() == '':
                continue
            if isinstance(v, str):
                try:
                    v = float(v.replace(' ', '').replace(',', '.'))
                except ValueError:
                    continue
            cells[(nm, gi)] = float(v)
    return cells


def main():
    path, label, period = sys.argv[1], sys.argv[2], sys.argv[3]
    cells = read_sheet(path)
    print('считано ячеек:', len(cells))
    if len(cells) < 3000:
        raise SystemExit('слишком мало ячеек — прерываю (ожидается ~3 800)')

    sys.path.insert(0, '/home/lnr/research-wiki-private/scripts')
    import db_tunnel
    conn = db_tunnel.connect()
    cur = conn.cursor()

    cur.execute("SELECT source_id FROM catalog.source WHERE source_code=%s",
                (SOURCE_CODE,))
    source_id = cur.fetchone()[0]

    cur.execute("SELECT release_id FROM catalog.release WHERE label=%s",
                (label,))
    row = cur.fetchone()
    if row:
        release_id = row[0]
        print('релиз уже есть:', release_id)
    else:
        cur.execute("""INSERT INTO catalog.release
            (source_id, label, published_at, status, notes)
            VALUES (%s, %s, %s, 'loaded', %s) RETURNING release_id""",
            (source_id, label, period,
             'ФНС 1-НОМ лист 1010 («всего»), графа 1-40, квартальный срез; '
             'агрегаты РФ/ФО с catalog.region.aggregation=source_aggregate'))
        release_id = cur.fetchone()[0]
    conn.commit()

    cur.execute("SELECT region_id, name_ru FROM catalog.region")
    by_name = {norm(n): r_ for r_, n in cur.fetchall()}
    cur.execute("SELECT metric_code, metric_id FROM catalog.metric "
                "WHERE metric_code ~ 'fns_1nom_[0-9]+$'")
    metrics = {c: i for c, i in cur.fetchall()}

    batch, unmapped = [], {}
    for (region_raw, col), v in cells.items():
        metric_id = metrics.get(f'fns_1nom_{col}')
        if metric_id is None:
            unmapped[f'col{col}'] = unmapped.get(f'col{col}', 0) + 1
            continue
        if region_raw.lower() == 'российская федерация':
            rid = 1
        else:
            rid = by_name.get(norm(region_raw))
        if rid is None:
            unmapped[region_raw] = unmapped.get(region_raw, 0) + 1
            continue
        batch.append((metric_id, rid, col))

    import io

    tmp = f'tmp_fns_{label}'
    cur.execute(f"""CREATE TEMP TABLE {tmp} (
        metric_id int, region_id int, period_start date,
        frequency_id int, value double precision,
        source_id int, release_id int) ON COMMIT DROP""")
    buf = io.StringIO()
    for (region_raw, col), v in cells.items():
        metric_id = metrics[f'fns_1nom_{col}']
        if region_raw.lower() == 'российская федерация':
            rid = 1
        else:
            rid = by_name.get(norm(region_raw))
        if rid is None or metric_id is None:
            continue
        buf.write("\t".join(map(str, (
            metric_id, rid, period, FREQ_YEAR, v,
            source_id, release_id))) + "\n")
    with cur.copy(f"COPY {tmp} FROM STDIN") as cp:
        cp.write(buf.getvalue())
    cur.execute(f"""INSERT INTO data.observation
        (metric_id, region_id, period_start, frequency_id, value,
         assessment_type, observation_status, source_id, release_id, sub_dimension)
        SELECT t.metric_id, t.region_id, t.period_start, t.frequency_id, t.value,
               'final', 'validated', t.source_id, t.release_id, ''
        FROM {tmp} t
        ON CONFLICT DO NOTHING""")

    cur.execute(f"""SELECT count(*) FROM data.observation
        WHERE release_id={release_id}""")
    print('в релизе %s: наблюдений %s' % (release_id, cur.fetchone()[0]))

    # приёмка: контрольные ячейки
    cur.execute("""SELECT o.value FROM data.observation o
        JOIN catalog.metric m USING(metric_id) WHERE m.metric_code='fns_1nom_26'
        AND o.release_id=%s AND o.region_id=1""", (release_id,))
    print('КТ транспортный налог РФ (гр26):', cur.fetchone())

    conn.commit()
    conn.close()
    print('готово; точка возврата: DELETE FROM data.observation '
          'WHERE release_id=%s' % release_id)


if __name__ == '__main__':
    main()
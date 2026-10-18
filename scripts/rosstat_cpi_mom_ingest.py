#!/usr/bin/env python3
"""Месячный ИПЦ Росстата (субъекты РФ) -> канон data.observation.

Источник: https://rosstat.gov.ru/storage/mediabank/ipc_RF_fo_sub_<MM>-<YYYY>.xlsx
Класс A (первоисточник; ЕМИСС 31074 публикует те же числа, но с адреса
датацентра отдаёт 403, поэтому канал — файл Росстата напрямую, без проверки
сертификата: у rosstat.gov.ru цепочка российского УЦ).

Метрика канона: emiss_31074_cpi_prevm_m (ИПЦ все товары и услуги,
% к предыдущему месяцу). На листе месяца раздел 1.1: колонка 0 — код ОКАТО
территории, колонка 1 — название, колонка 2 — ИПЦ м/м.

Питфоллы:
  * okato в core.region двузначный, поэтому сверка только по нормализованному
    имени плюс явные переопределения для веток Архангельска и Тюмени;
  * в файле «Архангельская область» — это полный агрегат (с НАО), а
    «Архангельская область (кроме Ненецкого автономного округа)» —
    остаточный ряд. В каноне наоборот по смыслу метки: (с НАО) = полный;
  * в названии ЯНАО встречается латинская «H» (гомоглиф) — нормализуем;
  * 4 новых субъекта в файле отсутствуют — они идут отдельной метрикой
    emiss_31074_cpi_prevm_rfnn_m.

Проверка перед записью: разбираем предыдущий месяц и сверяем с тем, что уже
лежит в каноне (август залит из ЕМИСС). Расхождение хотя бы по одной
территории > 0.005 останавливает запись, если не указан --force.

Запуск:
  ~/.hermes/hermes-agent/venv/bin/python3 scripts/rosstat_cpi_mom_ingest.py --dry-run
  ~/.hermes/hermes-agent/venv/bin/python3 scripts/rosstat_cpi_mom_ingest.py --apply
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import os
import re
import sys
from pathlib import Path

import openpyxl
import requests
import urllib3

sys.path.insert(0, str(Path(__file__).resolve().parent))
import db_tunnel  # noqa: E402

urllib3.disable_warnings()

METRIC_CODE = "emiss_31074_cpi_prevm_m"
SOURCE_ROSSTAT = 11
CACHE_DIR = Path(os.path.expanduser("~/personal/data/cpi-rosstat"))
URL_TPL = "https://rosstat.gov.ru/storage/mediabank/ipc_RF_fo_sub_{mm}-{yyyy}.xlsx"
TOL = 0.005

# Территории, где имя в файле не совпадает с меткой канона.
OVERRIDES = {
    "российская федерация": "ru",
    "архангельская область": "arkhangelsk_nao",
    "архангельская область (кроме ненецкого автономного округа)": "arkhangelsk",
    "ненецкий автономный округ": "nenets_ao",
    "ненецкий автономный округ ": "nenets_ao",
    "тюменская область": "tyumen_ao",
    "тюменская область (кроме ханты-мансийского автономного округа - югры "
    "и ямало-ненецкого автономного округа)": "tyumen",
    "ханты-мансийский автономный округ - югра": "khanty",
    "ямало-ненецкий автономный округ": "yamal",
}


def norm(name: str) -> str:
    """Нормализация названия: пробелы, переносы, гомоглифы, регистр."""
    s = (name or "").replace("\n", " ")
    s = s.replace("H", "Н").replace("O", "О").replace("P", "Р").replace("C", "С")
    s = re.sub(r"[\u00a0\s]+", " ", s).strip().lower()
    s = re.sub(r"^г\.\s*", "", s)          # «г. Москва» -> «москва»
    s = s.replace(" — ", " - ").replace(" – ", " - ")
    return s


def fetch_file(year: int, month: int) -> Path:
    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    path = CACHE_DIR / f"ipc_RF_fo_sub_{month:02d}-{year}.xlsx"
    if path.exists() and path.stat().st_size > 10_000:
        return path
    url = URL_TPL.format(mm=f"{month:02d}", yyyy=year)
    r = requests.get(url, timeout=60, verify=False,
                     headers={"User-Agent": "Mozilla/5.0"})
    if r.status_code != 200 or len(r.content) < 10_000:
        raise RuntimeError(f"не скачался {url}: HTTP {r.status_code}, "
                           f"{len(r.content)} байт")
    path.write_bytes(r.content)
    return path


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    h.update(path.read_bytes())
    return h.hexdigest()


def parse_month(path: Path, year: int, month: int) -> dict[str, tuple[str, float]]:
    """Лист 'MM(YYYY)' раздела 1.1 -> {нормализованное имя: (имя, ИПЦ м/м)}."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = f"{month:02d}({year})"
    if sheet not in wb.sheetnames:
        raise RuntimeError(f"в файле нет листа {sheet}; есть {wb.sheetnames[-3:]}")
    ws = wb[sheet]
    out: dict[str, tuple[str, float]] = {}
    started = False
    for row in ws.iter_rows(values_only=True):
        name = row[1] if len(row) > 1 else None
        val = row[2] if len(row) > 2 else None
        if name and str(name).strip().lower().startswith("наименование"):
            started = True          # шапка раздела 1.1
            continue
        if not started or not name:
            continue
        if isinstance(val, str):
            if val.strip().startswith("1)"):   # сноска внизу листа
                break
            try:
                val = float(val.replace(",", "."))
            except ValueError:
                continue
        if not isinstance(val, (int, float)):
            continue
        out[norm(str(name))] = (str(name).strip(), float(val))
    wb.close()
    if len(out) < 80:
        raise RuntimeError(f"лист {sheet}: разобрано только {len(out)} строк")
    return out


def build_mapping() -> dict[str, int]:
    """нормализованное имя РФ/территории -> region_id канона."""
    rows = db_tunnel.query(
        """select region_id, region_code, name_ru, short_name from core.region
           where is_active and level in ('country','region')""")
    m: dict[str, int] = {}
    by_code = {}
    for rid, code, name, short in rows:
        by_code[code] = rid
        if name:
            m.setdefault(norm(name), rid)
        if short:
            m.setdefault(norm(short), rid)
    for name, code in OVERRIDES.items():
        m[name] = by_code[code]
    return m


def canonical_prev(period: dt.date) -> dict[int, float]:
    rows = db_tunnel.query(
        """select o.region_id, o.value from data.observation o
           join catalog.metric m using(metric_id)
           where m.metric_code=%s and o.period_start=%s""",
        (METRIC_CODE, period))
    return {rid: float(v) for rid, v in rows}


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--year", type=int)
    ap.add_argument("--month", type=int, help="месяц данных (1-12)")
    ap.add_argument("--file", type=Path, help="локальный xlsx")
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--force", action="store_true",
                    help="писать даже при расхождении контрольной сверки")
    args = ap.parse_args()

    today = dt.date.today()
    year = args.year or today.year
    month = args.month or today.month

    path = args.file or fetch_file(year, month)
    target = parse_month(path, year, month)
    prev_year, prev_month = (year - 1, 12) if month == 1 else (year, month - 1)
    try:
        prev = parse_month(path, prev_year, prev_month)   # файл накопительный
    except RuntimeError:
        prev = parse_month(fetch_file(prev_year, prev_month), prev_year, prev_month)
    print(f"файл: {path}  sha256={sha256(path)[:16]}…")
    print(f"лист {month:02d}({year}): {len(target)} строк; "
          f"контрольный {prev_month:02d}({prev_year}): {len(prev)}")

    mapping = build_mapping()
    period = dt.date(year, month, 1)
    period_end = (dt.date(year + (month // 12), month % 12 + 1, 1)
                  - dt.timedelta(days=1))
    prev_period = dt.date(prev_year, prev_month, 1)

    # --- контроль: предыдущий месяц против канона ---
    canon_prev = canonical_prev(prev_period)
    checked = mismatched = 0
    bad: list[str] = []
    for name, (raw, val) in prev.items():
        rid = mapping.get(name)
        if rid is None or rid not in canon_prev:
            continue
        checked += 1
        if abs(canon_prev[rid] - val) > TOL:
            mismatched += 1
            bad.append(f"{raw}: файл {val} vs канон {canon_prev[rid]}")
    print(f"контроль {prev_period}: сверено {checked}, расхождений {mismatched}")
    for line in bad[:10]:
        print("   !", line)

    # --- сборка строк целевого месяца ---
    rows, unmapped = [], []
    for name, (raw, val) in target.items():
        rid = mapping.get(name)
        if rid is None:
            unmapped.append(raw)
            continue
        rows.append((rid, raw, val))
    print(f"к записи: {len(rows)} территорий; не сопоставлено: {len(unmapped)}")
    for u in unmapped:
        print("   ?", u)

    if (mismatched and not args.force):
        print("ОСТАНОВ: контрольная сверка не прошла (--force для обхода)")
        return 2

    already = db_tunnel.query(
        """select count(*) from data.observation o join catalog.metric m
           using(metric_id) where m.metric_code=%s and o.period_start=%s""",
        (METRIC_CODE, period))
    if already[0][0]:
        print(f"в каноне уже есть {already[0][0]} наблюдений за {period}")

    if not args.apply:
        sample = sorted(rows, key=lambda r: r[1])[:5]
        print("DRY-RUN. Примеры:", [(r[2], r[1]) for r in sample])
        rf = next((r for r in rows if r[0] == 1), None)
        print("РФ:", rf[2] if rf else "не найдена")
        return 0

    # --- запись: релиз + наблюдения ---
    label = f"rosstat_ipc_mom_{year}-{month:02d}"
    url = URL_TPL.format(mm=f"{month:02d}", yyyy=year)
    digest = sha256(path)
    rel = db_tunnel.query("select release_id from catalog.release where label=%s",
                          (label,))
    if rel:
        release_id = rel[0][0]
        db_tunnel.execute(
            """update catalog.release set status='loaded', ingested_at=now(),
               content_hash=%s, url=%s where release_id=%s""",
            (digest, url, release_id))
    else:
        release_id = db_tunnel.query(
            """insert into catalog.release
               (source_id, label, published_at, content_hash, status,
                ingested_at, url, notes)
               values (%s,%s,%s,%s,'loaded',now(),%s,%s) returning release_id""",
            (SOURCE_ROSSTAT, label, dt.datetime(year, month, 20, tzinfo=dt.timezone.utc),
             digest, url,
             "Региональный ИПЦ м/м, Росстат ipc_RF_fo_sub; класс A"))
        release_id = release_id[0][0]
    print(f"release_id={release_id}")

    db_tunnel.execute(
        """delete from data.observation where metric_id=
             (select metric_id from catalog.metric where metric_code=%s)
           and period_start=%s""", (METRIC_CODE, period))
    payload = [(rid, period, period_end, val, release_id, raw)
               for rid, raw, val in rows]
    n = db_tunnel.execute_many(
        """insert into data.observation
             (metric_id, region_id, period_start, period_end, value,
              observation_status, assessment_type, release_id, frequency_id,
              source_id, quality_flags, notes)
           values ((select metric_id from catalog.metric where metric_code=%s),
                   %s,%s,%s,%s,'validated','final',%s,5,%s,'{}', %s)""",
        [(METRIC_CODE, rid, ps, pe, val, release_id, SOURCE_ROSSTAT,
          f"Росстат ipc_RF_fo_sub_{month:02d}-{year}, {raw}")
         for rid, ps, pe, val, release_id, raw in payload])
    print(f"вставлено строк: {n}")

    check = db_tunnel.query(
        """select count(*), min(value), max(value) from data.observation o
           join catalog.metric m using(metric_id)
           where m.metric_code=%s and o.period_start=%s""",
        (METRIC_CODE, period))
    print(f"проверка: {check[0][0]} наблюдений за {period}, "
          f"min={check[0][1]}, max={check[0][2]}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

#!/usr/bin/env python3
"""Парсер официального файла СС-индексов Росстата sezon_2023_MM-YYYY.xlsx
в таблицу rosstat_ind_prod_saar (data/rosreestr_deals.db).

Колонки файла (лист '1'): Месяц | IPP (м/м факт, м/м СС, база факт, база СС) |
B Добыча | C Обрабатывающие | D Энергетика | E Водоснабжение — по 4 колонки на секцию.
Строка года маркером 'YYYY' / '20262' (сноска 2 — без ДНР/ЛНР/Запорожья/Херсона).

Запуск: python3 scripts/rosstat_sezon_parse.py <путь к xlsx> [год]
"""
import sys, sqlite3, datetime
from pathlib import Path
import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from yc_sync import ensure_db

MONTHS = {'январь': 1, 'февраль': 2, 'март': 3, 'апрель': 4, 'май': 5, 'июнь': 6,
          'июль': 7, 'август': 8, 'сентябрь': 9, 'октябрь': 10, 'ноябрь': 11, 'декабрь': 12}
SECTIONS = [('IPP', 1, 2, 3, 4), ('B_mining', 5, 6, 7, 8), ('C_manufacturing', 9, 10, 11, 12),
            ('D_energy', 13, 14, 15, 16), ('E_water', 17, 18, 19, 20)]

def main(path, only_year=None):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[1]] if wb.sheetnames[0] == 'Содержание' else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    year = None
    recs = []
    for r in rows:
        c0 = str(r[0]).strip() if r[0] is not None else ''
        if c0[:4].isdigit():
            year = int(c0[:4])
        if year and c0 in MONTHS:
            m = MONTHS[c0]
            for sec, i1, i2, i3, i4 in SECTIONS:
                recs.append((sec, f'{year}-{m:02d}-01', r[i1], r[i2], r[i3], r[i4]))
    if only_year:
        recs = [x for x in recs if x[1].startswith(str(only_year))]
    con = sqlite3.connect(str(ensure_db('rosreestr_deals.db')))
    updated = 0
    for sec, month, mf, ms, bf, bs in recs:
        old = con.execute("SELECT mom_fact, mom_saar, base_fact, base_saar FROM rosstat_ind_prod_saar WHERE section=? AND month=?",
                          (sec, month)).fetchone()
        new = tuple(None if v is None else float(v) for v in (mf, ms, bf, bs))
        if old is None or tuple(None if v is None else float(v) for v in old) != new:
            con.execute("INSERT OR REPLACE INTO rosstat_ind_prod_saar (section,month,mom_fact,mom_saar,base_fact,base_saar,source,ingested) VALUES (?,?,?,?,?,?,?,datetime('now'))",
                        (sec, month, *new, 'rosstat_sezon_2023'))
            updated += 1
            print('UPD', sec, month, 'old', old, '->', new)
    con.commit()
    print(f'Строк в файле: {len(recs)}, обновлено: {updated}')

if __name__ == '__main__':
    p = Path(sys.argv[1])
    if not p.exists():
        sys.exit(f'Файл не найден: {p}')
    main(p, sys.argv[2] if len(sys.argv) > 2 else None)
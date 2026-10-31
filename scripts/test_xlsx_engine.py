#!/usr/bin/env python3
"""Регрессия читателя xlsx: fastexcel против openpyxl на одном файле.

Проверяет то, что ломается при смене движка и не видно по времени прогона:
полную точность чисел в смешанных колонках, даты, ведущие нули в текстовых
кодах, пустые строки в начале листа и число строк.

Запуск: python3 scripts/test_xlsx_engine.py
"""
import datetime
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from xlsx_engine import cell_text, has_fastexcel, open_workbook  # noqa: E402


def _snapshot(path):
    out = {}
    for engine in ("openpyxl", "fastexcel"):
        wb = open_workbook(path, engine=engine)
        try:
            out[engine] = {
                name: [[cell_text(v) for v in r]
                       for r in wb[name].iter_rows(values_only=True)]
                for name in wb.sheetnames
            }
        finally:
            wb.close()
    return out


def _strip_tail(rows):
    return [r[:max((i + 1 for i, v in enumerate(r) if v), default=0)] for r in rows]


def main():
    import openpyxl
    from openpyxl import Workbook
    tmp = tempfile.mkdtemp(prefix="xlsx-engine-")
    path = os.path.join(tmp, "case.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "t"
    ws.append(["oktmo", "sum", "date", "mixed", "note"])
    ws.append(["00701000", 1234.56789012345, datetime.datetime(2025, 3, 1), "текст", None])
    ws.append(["0001", 0, datetime.datetime(2024, 12, 31), 42, "хвост"])
    ws.append([None, 1e-9, None, "0.1", " "])
    ws2 = wb.create_sheet("late")
    ws2["C4"] = "поздно"
    wb.save(path)

    snap = _snapshot(path)
    problems = []
    for sheet in snap["openpyxl"]:
        a = _strip_tail(snap["openpyxl"][sheet])
        b = _strip_tail(snap["fastexcel"][sheet])
        if a != b:
            problems.append((sheet, a, b))

    if not has_fastexcel():
        print("WARN: fastexcel не установлен — сверка вырождена")
    if problems:
        for sheet, a, b in problems:
            print(f"РАСХОЖДЕНИЕ на листе {sheet!r}:\n  openpyxl: {a}\n  fastexcel: {b}")
        return 1
    print(f"ок: значения совпадают на {len(snap['openpyxl'])} листах "
          f"(ведущие нули, даты, точность double, пустые строки сверху)")
    return 0


if __name__ == "__main__":
    sys.exit(main())

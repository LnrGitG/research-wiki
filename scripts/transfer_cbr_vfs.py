#!/usr/bin/env python3
"""Этап 1.4: ЦБ VFS 0420218 → catalog + data.observation.

Вход: /home/ubuntu/sandbox_datasets/raw_cbr_vfs/*.xlsx (7 файлов, срез 21.09).
Логика — 1:1 из loader/parse_cbr_vfs.py (SHEET_MAP, datekey, num, хеш строки).
Вставка executemany. Точка возврата: DELETE WHERE release_id=<новый>.
"""
import sys, os, io, datetime, hashlib, re, subprocess, unicodedata
sys.path.insert(0, os.path.dirname(__file__))
import db_tunnel

import sys as _sys
import datetime as _dt
RELEASE_LABEL = _sys.argv[1] if len(_sys.argv) > 1 else "cbr-vfs-0420218 2026-09-21 (перенос из v2, этап 1.4)"
try:
    _rd = _dt.date.fromisoformat(RELEASE_LABEL.split()[1])
except (ValueError, IndexError):
    _rd = _dt.date(2026, 9, 21)
SSH = ["ssh", "-o", "BatchMode=yes", "-i", "/home/lnr/.ssh/id_yc", "ubuntu@89.169.168.214"]
RAW_DIR = "/home/ubuntu/sandbox_datasets/raw_cbr_vfs/"

SHEET_MAP = {
    ("02_06","в рублях"): "cbr_izhk_count_m",
    ("02_06","итого"):    "cbr_izhk_count_m",
    ("02_07","в рублях"): "cbr_izhk_volume_m",
    ("02_07","итого"):    "cbr_izhk_volume_m",
    ("02_08","ставка в рублях"): "cbr_izhk_rate_m",
    ("02_08","срок в руб"):      "cbr_izhk_term_m",
    ("02_09","в рублях"): "cbr_izhk_debt_m",
    ("02_09","итого"):    "cbr_izhk_debt_m",
    ("02_15","в рублях"): "cbr_izhk_ddu_volume_m",
}
MON = {"Январь":1,"Февраль":2,"Март":3,"Апрель":4,"Май":5,"Июнь":6,"Июль":7,
       "Август":8,"Сентябрь":9,"Октябрь":10,"Ноябрь":11,"Декабрь":12}
ALIAS = {
    "архангельская область": "архангельская область (без нао)",
    "кемеровская область - кузбасс": "кемеровская область — кузбасс",
    "тюменская область": "тюменская область (без хмао и янао)",
}
UNIT_MAP = {1: 11, 4: 29, 5: 12, 8: 33}   # v2 unit → catalog (mln_rub 11, count 29, pct 12, months 33)
NAMES = {
    "cbr_izhk_count_m":      ("ЦБ: число выданных ипотечных кредитов (ИЖК), шт, за месяц", 4),
    "cbr_izhk_volume_m":     ("ЦБ: объём выдач ипотечных кредитов (ИЖК), млн руб., за месяц", 1),
    "cbr_izhk_rate_m":       ("ЦБ: средневзвешенная ставка по ипотечным кредитам (руб.), %", 5),
    "cbr_izhk_term_m":       ("ЦБ: средний срок кредитования по ИЖК (руб.), месяцев", 8),
    "cbr_izhk_debt_m":       ("ЦБ: задолженность по ипотечным кредитам (руб.), млн руб., на начало месяца", 1),
    "cbr_izhk_ddu_volume_m": ("ЦБ: объём выдач ипотечных кредитов на приобретение жилья с ДДУ, млн руб., за месяц", 1),
}


def datekey(v):
    if isinstance(v, datetime.datetime): return (v.year, v.month)
    if isinstance(v, str):
        m = re.match(r"(\S+)\s+(\d{4})", v.strip())
        if m and m.group(1) in MON: return (int(m.group(2)), MON[m.group(1)])
    return None


def num(x):
    if x is None: return None
    s = str(x).replace(" ","").replace(",",".")
    try: return float(s)
    except: return None


def main():
    import openpyxl
    names = subprocess.run(SSH + ["ls", RAW_DIR], capture_output=True, text=True).stdout.split()
    with db_tunnel.connect() as conn:
        with conn.cursor() as cur:
            src_id = cur.execute("SELECT source_id FROM catalog.source WHERE source_code='cbr'").fetchone()[0]
            freq_m = cur.execute("SELECT frequency_id FROM catalog.frequency WHERE freq_code='M'").fetchone()[0]
            rel = cur.execute("SELECT release_id FROM catalog.release WHERE label=%s",(RELEASE_LABEL,)).fetchone()
            release_id = rel[0] if rel else cur.execute("""INSERT INTO catalog.release
                (source_id,label,published_at,status,notes) VALUES (%s,%s,%s,'loaded',
                'перенос ЦБ VFS 0420218 (этап 1.4); сырьё ВМ xlsx, срез 2026-09-21') RETURNING release_id""",
                (src_id, RELEASE_LABEL, _rd)).fetchone()[0]
            conn.commit()
            region_rf = cur.execute("SELECT region_id FROM catalog.region WHERE region_code='ru'").fetchone()
            region_rf = region_rf[0] if region_rf else None
            rmap = {nm.strip().lower().replace("ё","е"): rid for rid, nm in
                    cur.execute("SELECT region_id, name_ru FROM catalog.region").fetchall()}
            rej_reg = set()
            # метрики
            midmap = {}
            for mc,(name,v2u) in NAMES.items():
                r = cur.execute("SELECT metric_id FROM catalog.metric WHERE metric_code=%s",(mc,)).fetchone()
                if r: midmap[mc]=r[0]
                else:
                    midmap[mc]=cur.execute("""INSERT INTO catalog.metric
                        (metric_code,name_ru,unit_id,source_id,frequency_id,metric_type,status,is_active)
                        VALUES (%s,%s,%s,%s,%s,'primary','active',true) RETURNING metric_id""",
                        (mc,name,UNIT_MAP[v2u],src_id,freq_m)).fetchone()[0]
                    conn.commit()
            rows=[]
            for f in sorted(names):
                if f[:5] not in ("02_04","02_05","02_06","02_07","02_08","02_09","02_15"): continue
                prefix = f[:5]
                raw = subprocess.run(SSH + ["cat", RAW_DIR+f], capture_output=True).stdout
                sha = hashlib.sha256(raw).hexdigest()
                wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
                for sn in wb.sheetnames:
                    key = (prefix, sn.strip().lower())
                    if key not in SHEET_MAP: continue
                    metric_code = SHEET_MAP[key]
                    grid = list(wb[sn].iter_rows(values_only=True))
                    hdr_i = next((i for i,r in enumerate(grid[:8]) if len([c for c in r[1:] if c is not None])>3), None)
                    if hdr_i is None: continue
                    hdr = grid[hdr_i]
                    n_file=0
                    for r_ in grid:
                        nm = str(r_[0]).strip() if r_[0] else ""
                        if not nm: continue
                        # первая колонка — имя региона (субъект по форме 0420218)
                        kk = nm.lower().replace("ё","е")
                        rid = region_rf if kk == "российская федерация" else rmap.get(kk) or (rmap.get(ALIAS[kk]) if kk in ALIAS else None)
                        # Агрегаты (РФ, ФО) помечены в catalog.region aggregation='source_aggregate'
                        # (не равны сумме субъектов — самостоятельные ряды источника).
                        if rid is None: rej_reg.add(nm); continue
                        for j in range(1, len(hdr)):
                            dk = datekey(hdr[j])
                            if dk is None: continue
                            val = num(r_[j])
                            if val is None: continue
                            rows.append((midmap[metric_code], rid, f"{dk[0]}-{dk[1]:02d}-01",
                                         freq_m, val, src_id, release_id))
                            n_file+=1
                    print(f"  {f}/{sn}: {n_file}", flush=True)
                wb.close()
            print("строк собрано:", len(rows))
            B=4000; inserted=0
            for i in range(0,len(rows),B):
                cur.executemany("""INSERT INTO data.observation
                    (metric_id,region_id,period_start,frequency_id,value,
                     assessment_type,observation_status,source_id,release_id,sub_dimension)
                    VALUES (%s,%s,%s,%s,%s,'final','validated',%s,%s,'')
                    ON CONFLICT (metric_id,region_id,period_start,frequency_id,
                                 assessment_type,sub_dimension,release_id) DO NOTHING""", rows[i:i+B])
                inserted += cur.rowcount
            conn.commit()
            print(f"вставлено: {inserted:,}; отказанные регионы: {sorted(rej_reg)[:8]} ({len(rej_reg)})")
            # сверка
            bad=[]
            for mc, mid in midmap.items():
                ref = cur.execute("""SELECT count(*) FROM v2.observation o JOIN v2.metric m
                    ON m.metric_id=o.metric_id WHERE m.metric_code=%s AND o.is_current""",(mc,)).fetchone()[0]
                got = cur.execute("SELECT count(*) FROM data.observation WHERE metric_id=%s",(mid,)).fetchone()[0]
                if got != ref: bad.append((mc, ref, got))
                print(f"  {mc:<22} canon={got:>6,} v2={ref:>6,} {'OK' if got==ref else 'РАСХ'}")

if __name__ == "__main__":
    main()